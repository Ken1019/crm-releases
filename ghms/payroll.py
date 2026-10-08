"""給与計算・給与明細・賃金台帳と、事業所の収支。

給与はタイムカード（実働・残業・深夜・夜勤回数）と職員の情報（給与の形・手当・保険の加入）から自動で計算し、
管理者が1項目ずつ直してから「確定」する。「職員に見せる」を押した明細だけ、本人の画面に出る（1人分ずつ）。

保険料率・所得税は概算。料率は毎年変わるので「給与の設定」で直す。所得税は国税庁の
「電子計算機等を使用して源泉徴収税額を計算する方法」（月額）で計算するが、正確な額は源泉徴収税額表で確認すること。
"""

import json
import math
from datetime import date

from flask import Blueprint, abort, flash, g, redirect, render_template, request, url_for
from openpyxl import Workbook

from . import excel
from .auth import admin_required, log_event
from .billing import compute_benefit
from .db import get_db, get_setting, now, set_setting
from .views import parse_date, parse_ym
from .work import month_summary, my_staff_id, setting_num

bp = Blueprint("payroll", __name__, url_prefix="/payroll")

# 給与の設定（キー, 名前, 初期値, 説明）
PAY_SETTINGS = [
    ("pay_monthly_hours", "月の所定労働時間（時間）", "160", "月給の人の残業代・深夜手当の時間単価に使います"),
    ("pay_ot_rate", "時間外の割増（%）", "25", "1日8時間をこえた分"),
    ("pay_night_rate", "深夜の割増（%）", "25", "22時〜翌5時"),
    ("pay_break_default", "休憩の目安（分）", "60", "6時間をこえる勤務で退勤するときの初期値"),
    ("pay_max_shift_hours", "1回の勤務の上限（時間）", "20", "これをこえると退勤を押せず「退勤忘れ」として管理者が直します（夜勤に合わせて）"),
    ("pay_fever", "体温のお知らせ（℃以上）", "37.5", ""),
    ("ins_health", "健康保険料率（%・労使合計）", "10.31", "協会けんぽ北海道 令和7年度の例。毎年3月に変わるので確認してください"),
    ("ins_care", "介護保険料率（%・労使合計）", "1.59", "40〜64歳の人だけ"),
    ("ins_pension", "厚生年金保険料率（%・労使合計）", "18.3", ""),
    ("ins_emp_ee", "雇用保険料率（%・本人）", "0.55", "令和7年度・一般の事業の例"),
    ("ins_emp_er", "雇用保険料率（%・事業所）", "0.9", ""),
    ("ins_child", "子ども・子育て拠出金率（%・事業所）", "0.36", ""),
    ("ins_rosai", "労災保険料率（%・事業所）", "0.3", "社会福祉施設の例"),
    ("pay_day", "支給日（明細に出す）", "翌月25日", ""),
]

EARNINGS = [("base", "基本給"), ("ot", "時間外手当"), ("night", "深夜手当"), ("yakin", "夜勤手当"), ("qual", "資格手当"),
            ("shogu", "処遇改善手当"), ("other", "その他手当"), ("commute", "交通費（非課税）")]
DEDUCTIONS = [("health", "健康保険"), ("care", "介護保険"), ("pension", "厚生年金"), ("emp", "雇用保険"),
              ("itax", "所得税"), ("rtax", "住民税"), ("other_ded", "その他控除")]
WORK_ITEMS = [("days", "出勤日数", "日"), ("hours", "実働時間", ""), ("ot_h", "時間外", ""), ("night_h", "深夜", ""), ("yakin_n", "夜勤", "回")]


def pset(key):
    return next(d for k, _, d, _ in PAY_SETTINGS if k == key)


def rate(key):
    return setting_num(key, pset(key)) / 100


def half_down(x):
    """社会保険料の本人負担：50銭以下切り捨て、50銭をこえたら切り上げ"""
    return max(0, math.ceil(x - 0.5))


def income_tax(amount, dependents, year):
    """源泉所得税（月額・甲欄）の概算。amount＝社会保険料等を引いた後の給与（交通費を除く）"""
    a = max(0, int(amount))
    if a < 1:
        return 0
    if year >= 2026:  # 令和8年1月以後（給与所得控除の最低額65万円・基礎控除58万円）
        if a <= 158333:
            ded = 54167
        elif a <= 299999:
            ded = math.ceil(a * 0.3 + 6667)
        elif a <= 549999:
            ded = math.ceil(a * 0.2 + 36667)
        elif a <= 708330:
            ded = math.ceil(a * 0.1 + 91667)
        else:
            ded = 162500
        basic = 48334 if a <= 2120833 else 40000 if a <= 2162499 else 26667 if a <= 2204166 else 13334 if a <= 2245833 else 0
    else:
        if a <= 135416:
            ded = 45834
        elif a <= 149999:
            ded = math.ceil(a * 0.4 - 8333)
        elif a <= 299999:
            ded = math.ceil(a * 0.3 + 6667)
        elif a <= 549999:
            ded = math.ceil(a * 0.2 + 36667)
        elif a <= 708330:
            ded = math.ceil(a * 0.1 + 91667)
        else:
            ded = 162500
        basic = 40000 if a <= 2162499 else 26667 if a <= 2204166 else 13334 if a <= 2245833 else 0
    c = a - ded - 31667 * max(0, int(dependents or 0)) - basic
    c = max(0, (c // 1000) * 1000) if c > 0 else 0
    if c <= 0:
        return 0
    for limit, r, sub in [(162500, 0.05105, 0), (275000, 0.1021, 8296), (579166, 0.2042, 36374), (750000, 0.23483, 54113),
                          (1500000, 0.33693, 130688), (3333333, 0.4084, 237893), (float("inf"), 0.45945, 408061)]:
        if c <= limit:
            tax = c * r - sub
            return max(0, int(round(tax / 10.0)) * 10)
    return 0


def _age_at(birth, d):
    b = parse_date(birth)
    if not b:
        return None
    return d.year - b.year - ((d.month, d.day) < (b.month, b.day))


def shogu_monthly(staff_id, first):
    fy = first.year if first.month >= 4 else first.year - 1
    row = get_db().execute("SELECT SUM(a.monthly) FROM shogu_allocations a JOIN shogu_plans p ON p.id=a.plan_id "
                           "WHERE a.staff_id=? AND p.fiscal_year=?", (staff_id, fy)).fetchone()
    return int(row[0] or 0)


def compute_pay(s, first, last, earnings=None):
    """1人1か月分。earnings を渡すと、その支給額から保険料・税を計算しなおす"""
    sm = month_summary(s["id"], first, last)
    hours, over_h, night_h = sm["total"] / 60, sm["over"] / 60, sm["night"] / 60
    pt = s["pay_type"] or "月給"
    base_salary, hourly, daily = s["base_salary"] or 0, s["hourly_wage"] or 0, s["daily_wage"] or 0
    if pt == "時給":
        unit = hourly
    elif pt == "日給":
        unit = daily / 8
    else:
        unit = (base_salary + (s["allowance_qual"] or 0)) / (setting_num("pay_monthly_hours", 160) or 160)
    if earnings is None:
        base = round(hourly * hours) if pt == "時給" else round(daily * sm["days"]) if pt == "日給" else int(base_salary)
        ot_factor = rate("pay_ot_rate") + (0 if pt == "時給" else 1)
        ctype = s["commute_type"] or "支給しない"
        commute = (s["commute"] or 0) if ctype == "毎月定額" else (s["commute"] or 0) * sm["days"] if ctype.startswith("1日") else 0
        earnings = {"base": base, "ot": round(unit * over_h * ot_factor), "night": round(unit * night_h * rate("pay_night_rate")),
                    "yakin": int((s["night_allowance"] or 0) * sm["yakin"]), "qual": int(s["allowance_qual"] or 0),
                    "shogu": shogu_monthly(s["id"], first), "other": int(s["allowance_other"] or 0), "commute": int(commute)}
    gross = sum(int(earnings.get(k) or 0) for k, _ in EARNINGS)
    commute = int(earnings.get("commute") or 0)
    ded = {k: 0 for k, _ in DEDUCTIONS}
    er = {"health": 0, "care": 0, "pension": 0, "emp": 0, "child": 0, "rosai": round((gross - commute) * rate("ins_rosai"))}
    if s["social_insurance"]:
        std = s["std_monthly"] or (gross - commute)
        ded["health"] = half_down(std * rate("ins_health") / 2)
        age = _age_at(s["birthdate"], last)
        if age is not None and 40 <= age < 65:
            ded["care"] = half_down(std * rate("ins_care") / 2)
        ded["pension"] = half_down(std * rate("ins_pension") / 2)
        er.update(health=ded["health"], care=ded["care"], pension=ded["pension"], child=round(std * rate("ins_child")))
    if s["employment_insurance"]:
        ded["emp"] = half_down(gross * rate("ins_emp_ee"))
        er["emp"] = round(gross * rate("ins_emp_er"))
    social = ded["health"] + ded["care"] + ded["pension"] + ded["emp"]
    ded["itax"] = income_tax(gross - commute - social, s["dependents"], first.year)
    ded["rtax"] = int(s["resident_tax"] or 0)
    total_ded = sum(ded.values())
    return {"earnings": {k: int(earnings.get(k) or 0) for k, _ in EARNINGS}, "deductions": ded,
            "work": {"days": sm["days"], "hours": sm["total"], "ot_h": sm["over"], "night_h": sm["night"], "yakin_n": sm["yakin"]},
            "missing": sm["missing"], "gross": gross, "total_ded": total_ded, "net": gross - total_ded,
            "employer": er, "employer_total": sum(er.values()), "pay_type": pt, "unit": round(unit)}


def _totals(data):
    data["gross"] = sum(int(v or 0) for v in data["earnings"].values())
    data["total_ded"] = sum(int(v or 0) for v in data["deductions"].values())
    data["net"] = data["gross"] - data["total_ded"]
    return data


def load_slip(staff_id, ym):
    row = get_db().execute("SELECT * FROM payslips WHERE staff_id=? AND ym=?", (staff_id, ym)).fetchone()
    if not row:
        return None
    data = json.loads(row["data"])
    data.update(status=row["status"], updated_by=row["updated_by"], updated_at=row["updated_at"])
    return data


def save_slip(staff_id, ym, data, status):
    _totals(data)
    keep = {k: data[k] for k in ("earnings", "deductions", "work", "gross", "total_ded", "net", "employer", "employer_total",
                                 "pay_type", "unit", "shared", "memo") if k in data}
    get_db().execute("INSERT OR REPLACE INTO payslips (staff_id, ym, data, gross, deductions, net, status, updated_by, updated_at)"
                     " VALUES (?,?,?,?,?,?,?,?,?)", (staff_id, ym, json.dumps(keep, ensure_ascii=False), data["gross"],
                                                     data["total_ded"], data["net"], status, g.user["username"], now()))


def payroll_staff(first, last):
    """その月に在籍していた職員（退職者もその月までの分は出す）"""
    return get_db().execute("SELECT * FROM staff WHERE (status IS NULL OR status != '退職' OR id IN "
                            "(SELECT staff_id FROM timecards WHERE date BETWEEN ? AND ?)) ORDER BY kana, name",
                            (first.isoformat(), last.isoformat())).fetchall()


def month_payroll(first, last):
    ym = first.strftime("%Y-%m")
    rows = []
    for s in payroll_staff(first, last):
        slip = load_slip(s["id"], ym)
        rows.append({"s": s, "d": slip or compute_pay(s, first, last), "saved": slip is not None})
    return rows


def hm(minutes):
    return f"{int(minutes) // 60}:{int(minutes) % 60:02d}" if minutes else "0:00"


@bp.route("/", methods=["GET", "POST"])
@admin_required
def index():
    first, last = parse_ym(request.values.get("ym"))
    ym = first.strftime("%Y-%m")
    rows = month_payroll(first, last)
    if request.method == "POST" and request.form.get("action") == "save_all":
        n = 0
        for r in rows:
            if not r["saved"]:
                save_slip(r["s"]["id"], ym, r["d"], "下書き")
                n += 1
        log_event("payroll_save", "payslips", None, f"{ym} {n}人")
        get_db().commit()
        flash(f"{first:%Y年%m月}の給与を{n}人分、下書きとして保存しました。内容を確かめて「確定」してください。", "ok")
        return redirect(url_for("payroll.index", ym=ym))
    total = {k: sum(r["d"][k] for r in rows) for k in ("gross", "total_ded", "net", "employer_total")}
    return render_template("payroll.html", rows=rows, ym=ym, first=first, total=total, hm=hm, EARNINGS=EARNINGS)


@bp.route("/<int:sid>/<ym>", methods=["GET", "POST"])
@admin_required
def edit(sid, ym):
    db = get_db()
    s = db.execute("SELECT * FROM staff WHERE id=?", (sid,)).fetchone() or abort(404)
    first, last = parse_ym(ym)
    ym = first.strftime("%Y-%m")
    slip = load_slip(sid, ym)
    auto = compute_pay(s, first, last)
    if request.method == "POST":
        action = request.form.get("action")
        if slip and slip["status"] == "確定" and action not in ("unlock", "share", "unshare"):
            flash("確定した給与は直せません。直すときは「確定を取り消す」を押してください。", "error")
            return redirect(url_for("payroll.edit", sid=sid, ym=ym))
        if action == "reset":
            db.execute("DELETE FROM payslips WHERE staff_id=? AND ym=?", (sid, ym))
            flash("自動計算に戻しました。", "ok")
        elif action in ("unlock", "share", "unshare"):
            if not slip:
                abort(400)
            if action == "unlock":
                save_slip(sid, ym, dict(slip, shared=False), "下書き")
                flash("確定を取り消しました（職員の画面からも見えなくなります）。", "ok")
            else:
                if slip["status"] != "確定":
                    flash("職員に見せるのは、確定してからです。", "error")
                    return redirect(url_for("payroll.edit", sid=sid, ym=ym))
                save_slip(sid, ym, dict(slip, shared=action == "share"), "確定")
                flash("職員の画面に明細を出しました。" if action == "share" else "職員の画面から明細を外しました。", "ok")
            log_event("payroll_" + action, "payslips", sid, ym)
        else:
            earnings = {k: request.form.get(f"e_{k}", type=int) or 0 for k, _ in EARNINGS}
            if action == "recalc":
                data = compute_pay(s, first, last, earnings)
                flash("支給額から保険料・税を計算しなおしました（まだ保存していません。下の「保存する」を押してください）。", "ok")
                return render_template("payroll_edit.html", s=s, ym=ym, first=first, d=data, auto=auto, slip=slip,
                                       EARNINGS=EARNINGS, DEDUCTIONS=DEDUCTIONS, WORK_ITEMS=WORK_ITEMS, hm=hm, unsaved=True)
            base = slip or auto
            data = dict(base, earnings=earnings, deductions={k: request.form.get(f"d_{k}", type=int) or 0 for k, _ in DEDUCTIONS},
                        memo=request.form.get("memo", ""), shared=False)
            status = "確定" if action == "confirm" else "下書き"
            save_slip(sid, ym, data, status)
            log_event("payroll_" + ("confirm" if status == "確定" else "edit"), "payslips", sid, ym)
            flash(f"{s['name']} さんの{first:%Y年%m月}の給与を{'確定' if status == '確定' else '保存'}しました。", "ok")
        db.commit()
        return redirect(url_for("payroll.edit", sid=sid, ym=ym))
    return render_template("payroll_edit.html", s=s, ym=ym, first=first, d=slip or auto, auto=auto, slip=slip,
                           EARNINGS=EARNINGS, DEDUCTIONS=DEDUCTIONS, WORK_ITEMS=WORK_ITEMS, hm=hm, unsaved=False)


def _slip_view(s, ym, d):
    first, _ = parse_ym(ym)
    return render_template("payroll_slip.html", s=s, first=first, d=d, EARNINGS=EARNINGS, DEDUCTIONS=DEDUCTIONS,
                           WORK_ITEMS=WORK_ITEMS, hm=hm, pay_day=get_setting("pay_day", pset("pay_day")),
                           office={"office_name": get_setting("office_name"), "corp_name": get_setting("corp_name")})


@bp.route("/<int:sid>/<ym>/slip")
@admin_required
def slip(sid, ym):
    s = get_db().execute("SELECT * FROM staff WHERE id=?", (sid,)).fetchone() or abort(404)
    first, last = parse_ym(ym)
    d = load_slip(sid, first.strftime("%Y-%m")) or compute_pay(s, first, last)
    return _slip_view(s, first.strftime("%Y-%m"), d)


@bp.route("/mine")
def mine():
    """職員：自分の給与明細（管理者が「職員に見せる」を押した月だけ）"""
    sid = my_staff_id()
    rows = []
    if sid:
        for r in get_db().execute("SELECT * FROM payslips WHERE staff_id=? AND status='確定' ORDER BY ym DESC", (sid,)):
            if json.loads(r["data"]).get("shared"):
                rows.append(r)
    return render_template("payroll_mine.html", rows=rows, linked=bool(sid))


@bp.route("/mine/<ym>")
def mine_slip(ym):
    sid = my_staff_id()
    d = load_slip(sid, ym) if sid else None
    if not d or d["status"] != "確定" or not d.get("shared"):
        abort(404)
    s = get_db().execute("SELECT * FROM staff WHERE id=?", (sid,)).fetchone()
    log_event("view", "payslips", sid, f"{ym} 本人")
    get_db().commit()
    return _slip_view(s, ym, d)


@bp.route("/ledger.xlsx")
@admin_required
def ledger():
    """賃金台帳（1年分・職員ごとに1枚）"""
    year = request.args.get("year", type=int) or date.today().year
    wb = Workbook()
    wb.remove(wb.active)
    for s in get_db().execute("SELECT * FROM staff ORDER BY kana, name"):
        months = []
        for m in range(1, 13):
            first, last = parse_ym(f"{year}-{m:02d}")
            months.append(load_slip(s["id"], f"{year}-{m:02d}"))
        if not any(months):
            continue
        rows = []
        for group, items in (("work", WORK_ITEMS), ("earnings", EARNINGS), ("deductions", DEDUCTIONS)):
            for k, label, *unit in items:
                vals = [((x or {}).get(group) or {}).get(k, 0) or 0 for x in months]
                if group == "work" and k in ("hours", "ot_h", "night_h"):
                    rows.append([label] + [hm(v) for v in vals] + [hm(sum(vals))])
                else:
                    rows.append([label] + vals + [sum(vals)])
            if group == "earnings":
                vals = [(x or {}).get("gross", 0) for x in months]
                rows.append(["総支給額"] + vals + [sum(vals)])
        vals = [(x or {}).get("total_ded", 0) for x in months]
        rows.append(["控除合計"] + vals + [sum(vals)])
        vals = [(x or {}).get("net", 0) for x in months]
        rows.append(["差引支給額"] + vals + [sum(vals)])
        excel.add_table(wb.create_sheet(excel.safe_sheet_title(s["name"])), f"賃金台帳　{year}年　{s['name']}",
                        ["項目"] + [f"{m}月" for m in range(1, 13)] + ["合計"], rows,
                        subtitle=f"入職日 {s['hire_date'] or ''}　職種 {s['job'] or ''}　雇用形態 {s['employment'] or ''}（確定・下書きで保存した月）",
                        widths=[14] + [10] * 13)
    if not wb.sheetnames:
        wb.create_sheet("なし")
    return excel.send_workbook(wb, f"賃金台帳_{year}")


@bp.route("/settings", methods=["GET", "POST"])
@admin_required
def settings():
    if request.method == "POST":
        for k, _, d, _ in PAY_SETTINGS:
            v = (request.form.get(k) or "").strip()
            set_setting(k, v or d)
        set_setting("pay_health_required", "1" if request.form.get("pay_health_required") else "0")
        log_event("settings", "payroll", None, "給与の設定")
        get_db().commit()
        flash("給与の設定を保存しました。", "ok")
        return redirect(url_for("payroll.settings"))
    values = {k: get_setting(k, d) for k, _, d, _ in PAY_SETTINGS}
    return render_template("payroll_settings.html", PAY_SETTINGS=PAY_SETTINGS, values=values,
                           health_required=get_setting("pay_health_required", "1") == "1")


# ---------------------------------------------------------------- 事業所の収支
INVOICE_ITEMS = [("rent", "家賃"), ("food", "食費"), ("utility", "光熱水費"), ("daily_goods", "日用品費"), ("other_amount", "その他")]


def month_profit(first, last):
    db = get_db()
    ym = first.strftime("%Y-%m")
    results, _, _ = compute_benefit(first, last)
    inc = {"給付費（国保連・概算）": sum(x["yen"] - x["burden"] for x in results),
           "利用者負担": sum(x["burden"] for x in results)}
    inv = db.execute("SELECT " + ", ".join(f"SUM({k})" for k, _ in INVOICE_ITEMS) + " FROM invoices WHERE ym=?", (ym,)).fetchone()
    for (k, label), v in zip(INVOICE_ITEMS, inv):
        inc[label] = int(v or 0)
    pay = month_payroll(first, last)
    out = {"給与（総支給）": sum(r["d"]["gross"] for r in pay), "法定福利費（事業所負担の保険料）": sum(r["d"]["employer_total"] for r in pay)}
    for row in db.execute("SELECT kind, SUM(amount) FROM expenses WHERE date BETWEEN ? AND ? GROUP BY kind",
                          (first.isoformat(), last.isoformat())):
        out[row[0] or "その他"] = int(row[1] or 0)
    income, expense = sum(inc.values()), sum(out.values())
    labor = out["給与（総支給）"] + out["法定福利費（事業所負担の保険料）"]
    return {"first": first, "income": inc, "expense": out, "income_total": income, "expense_total": expense,
            "profit": income - expense, "labor": labor, "labor_ratio": round(labor * 100 / income, 1) if income else None,
            "estimated": not all(r["saved"] for r in pay) or not any(inv)}


@bp.route("/profit")
@admin_required
def profit():
    today = date.today()
    fy = request.args.get("fy", type=int) or (today.year if today.month >= 4 else today.year - 1)
    months = []
    for i in range(12):
        y, m = (fy, 4 + i) if i < 9 else (fy + 1, i - 8)
        first, last = parse_ym(f"{y}-{m:02d}")
        if first > today.replace(day=1):
            break
        months.append(month_profit(first, last))
    inc_keys = list(dict.fromkeys(k for x in months for k in x["income"]))
    exp_keys = list(dict.fromkeys(k for x in months for k in x["expense"]))
    return render_template("payroll_profit.html", fy=fy, months=months, inc_keys=inc_keys, exp_keys=exp_keys,
                           inc_tot={k: sum(x["income"].get(k, 0) for x in months) for k in inc_keys},
                           exp_tot={k: sum(x["expense"].get(k, 0) for x in months) for k in exp_keys},
                           total={k: sum(x[k] for x in months) for k in ("income_total", "expense_total", "profit", "labor")})
