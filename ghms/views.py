import calendar
import io
import json
import sqlite3
from datetime import date, datetime, timedelta

from flask import Blueprint, abort, current_app, flash, g, redirect, render_template, request, send_file, url_for
from openpyxl import Workbook

from . import excel
from .auth import admin_required
from .crud import audit, label_of, ref_options
from .customize import feature_on, tracked_meetings
from .db import get_db, get_setting, now, set_setting
from .forms import NUM_SETTINGS, check_setting, db_int, safe_float, safe_int, setting_number
from .entities import MEAL, MED, MOOD, TIME_SLOT
from .hubs import HUB_BY_KEY, visible_hubs, visible_tasks

bp = Blueprint("views", __name__)

ACTIVE_RES = "status IS NULL OR status != '退居'"
# (種別, この日数を過ぎたら注意)
# ホームに最終実施日を出す会議は「設定」→「選択肢を変える」で決める（customize.tracked_meetings）


# 日付として受けつける年（9999年・1年などは、前後の日を計算するとエラーになるため使わない）
DATE_YEARS = (1900, 2100)
YM_YEARS = (2000, 2100)


def parse_date(s, default=None):
    try:
        d = datetime.strptime(s, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return default
    return d if DATE_YEARS[0] <= d.year <= DATE_YEARS[1] else default


def parse_ym(s):
    try:
        d = datetime.strptime(s, "%Y-%m").date()
    except (TypeError, ValueError):
        d = None
    if d is None or not YM_YEARS[0] <= d.year <= YM_YEARS[1]:
        d = date.today().replace(day=1)
    last = date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])
    return d, last


def fiscal_year(d):
    return d.year if d.month >= 4 else d.year - 1


def years_between(start, end):
    if not start:
        return None
    y = end.year - start.year - ((end.month, end.day) < (start.month, start.day))
    return max(y, 0)


# ---------------------------------------------------------------- ダッシュボード
def _alerts_and_away():
    """受給者証・計画などの期限のお知らせと、いま不在の入居者"""
    db = get_db()
    today = date.today()
    # まちがった値が保存されていても、ホームが開けなくならないように（初期値・範囲の中におさめる）
    cert_days = setting_number("cert_alert_days", 60)
    plan_days = setting_number("plan_alert_days", 30)
    cert_limit = (today + timedelta(days=cert_days)).isoformat()
    plan_limit = (today + timedelta(days=plan_days)).isoformat()

    alerts = []
    for r in db.execute(f"SELECT * FROM residents WHERE ({ACTIVE_RES}) AND cert_end IS NOT NULL AND cert_end <= ? "
                        "ORDER BY cert_end", (cert_limit,)):
        alerts.append(("受給者証（支給決定期間）", r["name"], r["cert_end"], url_for("crud.view", key="residents", rid=r["id"])))
    for r in db.execute(f"SELECT * FROM residents WHERE ({ACTIVE_RES}) AND level_end IS NOT NULL AND level_end <= ? "
                        "ORDER BY level_end", (cert_limit,)):
        alerts.append(("障害支援区分の有効期限", r["name"], r["level_end"], url_for("crud.view", key="residents", rid=r["id"])))
    for p in db.execute("SELECT p.*, r.name AS rname FROM support_plans p JOIN residents r ON r.id=p.resident_id "
                        "WHERE (r.status IS NULL OR r.status != '退居') AND (p.status IS NULL OR p.status != '終了') "
                        "AND p.period_end IS NOT NULL AND p.period_end <= ? ORDER BY p.period_end", (plan_limit,)):
        alerts.append(("個別支援計画の期間終了", p["rname"], p["period_end"], url_for("crud.view", key="support_plans", rid=p["id"])))
    plans = db.execute("SELECT p.*, r.name AS rname FROM support_plans p JOIN residents r ON r.id=p.resident_id "
                       "WHERE (r.status IS NULL OR r.status != '退居') AND (p.status IS NULL OR p.status != '終了')").fetchall()
    for p in sorted((p for p in plans if p["next_monitoring"] and p["next_monitoring"] <= plan_limit),
                    key=lambda p: p["next_monitoring"]):
        alerts.append(("モニタリング予定", p["rname"], p["next_monitoring"], url_for("crud.view", key="support_plans", rid=p["id"])))
    # 個別支援計画は少なくとも6か月に1回見直す（モニタリング）。作成日・同意日・モニタリング済にした日のうち新しいものから数える
    for p in plans:
        if p["status"] == "作成中":
            continue
        base = [d for d in (parse_date(p["created_on"]), parse_date(p["consent_date"]),
                            parse_date((p["updated_at"] or "")[:10]) if p["status"] == "モニタリング済" else None) if d]
        if not base:
            continue
        due = (max(base) + timedelta(days=183)).isoformat()
        if due <= plan_limit and not (p["next_monitoring"] and p["next_monitoring"] <= plan_limit):
            alerts.append(("個別支援計画の見直し（6か月）", p["rname"], due, url_for("crud.view", key="support_plans", rid=p["id"])))
    # 入居中なのに有効な個別支援計画がない
    for r in db.execute(f"SELECT * FROM residents WHERE ({ACTIVE_RES}) AND id NOT IN "
                        "(SELECT resident_id FROM support_plans WHERE status IS NULL OR status != '終了')"):
        alerts.append(("個別支援計画が未作成", r["name"], "", url_for("crud.new", key="support_plans", resident_id=r["id"])))
    for d in ([] if not feature_on("documents") else db.execute("SELECT d.*, r.name AS rname FROM resident_documents d JOIN residents r ON r.id=d.resident_id "
                        f"WHERE (r.{ACTIVE_RES.replace(' OR status', ' OR r.status')}) AND d.expires_on IS NOT NULL AND d.expires_on <= ?",
                        (cert_limit,))):
        alerts.append((f"{d['doc_type']}の更新", d["rname"], d["expires_on"], url_for("crud.view", key="resident_documents", rid=d["id"])))
    if g.user["role"] == "admin" and feature_on("invoices"):
        for i in db.execute("SELECT i.*, r.name AS rname FROM invoices i JOIN residents r ON r.id=i.resident_id "
                            "WHERE i.status != '入金済' AND i.due_date IS NOT NULL AND i.due_date < ?", (today.isoformat(),)):
            alerts.append((f"{i['ym']}分の利用料が未入金", i["rname"], i["due_date"], url_for("crud.edit", key="invoices", rid=i["id"])))
    from .absences import current_absences, sync_open

    sync_open()
    away = current_absences() if feature_on("absences") else []
    for x in away:
        if x["need_contact"]:
            alerts.append((f"入院中の連絡が{x['days']}日ありません", x["a"]["rname"], "",
                           url_for("absences.detail", aid=x["a"]["id"])))
    alerts.sort(key=lambda a: a[2] or "0000")
    return alerts, away


def dashboard_alerts():
    return _alerts_and_away()[0]


@bp.route("/")
def dashboard():
    db = get_db()
    today = date.today()
    alerts, away = _alerts_and_away()

    meetings = []
    if feature_on("meetings"):
        lasts = dict(db.execute("SELECT kind, MAX(date) FROM meetings GROUP BY kind").fetchall())
        for kind, limit in tracked_meetings():
            last = lasts.get(kind)
            meetings.append((kind, last, (not last) or parse_date(last, today) < today - timedelta(days=limit)))

    homes = db.execute("SELECT * FROM homes ORDER BY name").fetchall()
    logged = {r[0] for r in db.execute("SELECT home_id FROM daily_logs WHERE date=?", (today.isoformat(),))}
    counts = dict(db.execute(f"SELECT home_id, COUNT(*) FROM residents WHERE {ACTIVE_RES} GROUP BY home_id").fetchall())
    home_stats = [(h, counts.get(h["id"], 0), h["id"] in logged) for h in homes]

    stats = {
        "residents": db.execute(f"SELECT COUNT(*) FROM residents WHERE {ACTIVE_RES}").fetchone()[0],
        "staff": db.execute("SELECT COUNT(*) FROM staff WHERE status IS NULL OR status='在籍'").fetchone()[0],
        "incidents": db.execute("SELECT COUNT(*) FROM incidents WHERE date >= ?",
                                ((today - timedelta(days=30)).isoformat(),)).fetchone()[0],
        "addons": db.execute("SELECT COUNT(*) FROM addons WHERE active=1").fetchone()[0],
    }
    recent = db.execute("SELECT i.*, r.name AS rname FROM incidents i LEFT JOIN residents r ON r.id=i.resident_id "
                        "ORDER BY i.date DESC, i.id DESC LIMIT 5").fetchall()
    update_available = None
    if g.user["role"] == "admin":
        from .updater import cached_latest, maybe_check_in_background

        maybe_check_in_background(current_app._get_current_object())
        update_available = cached_latest()
    birthdays = []
    for r in db.execute(f"SELECT id, name, birthdate FROM residents WHERE ({ACTIVE_RES}) AND birthdate LIKE ?",
                        (f"%-{today.month:02d}-%",)):
        b = parse_date(r["birthdate"])
        if b and b.month == today.month:
            birthdays.append({"id": r["id"], "name": r["name"], "day": b.day, "age": today.year - b.year, "today": b.day == today.day})
    birthdays.sort(key=lambda x: x["day"])
    all_tasks = [dict(t, hub=h) for h in visible_hubs() for t in visible_tasks(h)]
    from .today import build

    from flask import session

    pin_admin = g.user["role"] == "admin" and session.get("via") == "pin"
    if g.user["role"] != "admin" or pin_admin:
        # 職員の画面：打刻・今日の記録・自分に関係するお知らせだけ。
        # PINで入った管理者も、お金や職員の体調はパスワードで本人確認してから（下の「管理者の画面」から）
        from .work import my_staff_id, my_upcoming_shifts, open_card

        sid = my_staff_id()
        return render_template("dashboard_staff.html", todo=build(user_admin=False, staff_id=sid), pin_admin=pin_admin, birthdays=birthdays, all_tasks=all_tasks, away=away, home_stats=home_stats,
                               today=today, card=open_card(sid) if sid and feature_on("timecard") else None, sid=sid,
                               shifts=my_upcoming_shifts(sid) if sid and feature_on("shift") else [])
    work_today = profit_now = None
    if feature_on("timecard"):
        from .work import today_status

        work_today = today_status()
    if feature_on("payroll"):
        from .payroll import month_profit

        first = today.replace(day=1)
        profit_now = month_profit(first, (first + timedelta(days=32)).replace(day=1) - timedelta(days=1))
    return render_template("dashboard.html", todo=build(user_admin=True, alerts=alerts), work_today=work_today, profit_now=profit_now, birthdays=birthdays, update_available=update_available, all_tasks=all_tasks, away=away, alerts=alerts, meetings=meetings, home_stats=home_stats, stats=stats,
                           recent=recent, today=today)


# ---------------------------------------------------------------- 目的別メニュー
@bp.route("/do/<key>")
def hub(key):
    h = HUB_BY_KEY.get(key)
    if h is None or h not in visible_hubs():
        abort(404)
    return render_template("hub.html", hub=h, tasks=visible_tasks(h))


# ---------------------------------------------------------------- 日誌（1日分まとめて入力）
LOG_FIELDS = ["day_staff", "night_staff", "residents_count", "absent", "summary", "events", "handover", "checker"]
REC_FIELDS = ["temperature", "meal", "medication", "mood", "content", "staff"]


def _home_residents(home_id):
    if home_id:
        return get_db().execute(f"SELECT * FROM residents WHERE home_id=? AND ({ACTIVE_RES}) ORDER BY room, kana",
                                (home_id,)).fetchall()
    return get_db().execute(f"SELECT * FROM residents WHERE {ACTIVE_RES} ORDER BY room, kana").fetchall()


@bp.route("/journal", methods=["GET", "POST"])
def journal():
    db = get_db()
    d = parse_date(request.values.get("date"), date.today())
    homes = db.execute("SELECT * FROM homes ORDER BY name").fetchall()
    home_id = request.values.get("home_id", type=db_int)
    if home_id is None:
        home_id = homes[0]["id"] if homes else None
    # home_id=0 は「住居未設定」（住居が入っていない入居者）。業務日誌は住居なし（NULL）として保存する
    log_home = None if home_id == 0 else home_id
    slot = request.values.get("slot") if request.values.get("slot") in TIME_SLOT else "終日"
    if home_id == 0:
        residents = db.execute(f"SELECT * FROM residents WHERE home_id IS NULL AND ({ACTIVE_RES}) ORDER BY room, kana").fetchall()
    else:
        residents = _home_residents(home_id)

    if request.method == "POST":
        log = {f: (request.form.get(f, "").strip() or None) for f in LOG_FIELDS}
        if log["residents_count"]:
            log["residents_count"] = safe_int(log["residents_count"], None, 0, 100000)
        existing = db.execute("SELECT id FROM daily_logs WHERE date=? AND home_id IS ?", (d.isoformat(), log_home)).fetchone()
        cols = list(log)
        if existing:
            # すでにある日誌は、全部消したときも消した内容で上書きする
            db.execute(f"UPDATE daily_logs SET {', '.join(c + '=?' for c in cols)}, updated_at=?, updated_by=? WHERE id=?",
                       [log[c] for c in cols] + [now(), g.user["username"], existing["id"]])
            audit("update", "daily_logs", existing["id"])
        elif any(v not in (None, "") for v in log.values()):
            cur = db.execute(f"INSERT INTO daily_logs (date, home_id, {', '.join(cols)}, created_at, updated_at, updated_by)"
                             f" VALUES (?, ?, {', '.join('?' for _ in cols)}, ?, ?, ?)",
                             [d.isoformat(), log_home] + [log[c] for c in cols] + [now(), now(), g.user["username"]])
            audit("create", "daily_logs", cur.lastrowid)
        saved = removed = 0
        by_id = request.form.get("by_id") == "1"  # 今の画面は、行ごとに記録の番号を送る
        for r in residents:
            rec = {f: (request.form.get(f"r{r['id']}_{f}", "").strip() or None) for f in REC_FIELDS}
            if rec["temperature"]:
                rec["temperature"] = safe_float(rec["temperature"], None, 20, 50)
            # 画面に出していた記録（行ごとの番号）だけを直す。同じ日・時間帯のほかの記録は上書きしない
            rec_id = request.form.get(f"r{r['id']}_id", type=db_int)
            if rec_id:
                ex = db.execute("SELECT id FROM support_records WHERE id=? AND resident_id=?", (rec_id, r["id"])).fetchone()
            elif by_id:
                ex = None  # 画面を開いたときに記録がなかった → 新しく作る（あとから書かれた記録は上書きしない）
            else:
                # 番号のない送信（古い画面など）は、同じ日・時間帯のいちばん新しい記録を直す
                ex = db.execute("SELECT id FROM support_records WHERE date=? AND resident_id=? AND time_slot=? ORDER BY id DESC",
                                (d.isoformat(), r["id"], slot)).fetchone()
            # 記録者欄は自動入力されるため、それ以外に記入がある場合のみ記録として扱う
            filled = any(rec[f] not in (None, "") for f in REC_FIELDS if f != "staff")
            cols = list(rec)
            if ex and not filled and not (rec_id or by_id):
                continue  # 番号のない送信で空欄のときは、記録を変えない
            if ex and not filled:
                # 記録の欄を全部消したら、その記録を削除する
                db.execute("DELETE FROM support_records WHERE id=?", (ex["id"],))
                audit("delete", "support_records", ex["id"])
                removed += 1
            elif ex:
                db.execute(f"UPDATE support_records SET {', '.join(c + '=?' for c in cols)}, updated_at=?, updated_by=? WHERE id=?",
                           [rec[c] for c in cols] + [now(), g.user["username"], ex["id"]])
                audit("update", "support_records", ex["id"])
                saved += 1
            elif filled:
                cur = db.execute(f"INSERT INTO support_records (date, resident_id, time_slot, {', '.join(cols)}, created_at, updated_at, updated_by)"
                                 f" VALUES (?, ?, ?, {', '.join('?' for _ in cols)}, ?, ?, ?)",
                                 [d.isoformat(), r["id"], slot] + [rec[c] for c in cols] + [now(), now(), g.user["username"]])
                audit("create", "support_records", cur.lastrowid)
                saved += 1
        db.commit()
        flash(f"保存しました（支援記録 {saved}件" + (f"・削除 {removed}件" if removed else "") + "）。", "ok")
        return redirect(url_for("views.journal", date=d.isoformat(), home_id=home_id, slot=slot))

    from .absences import current_absences

    away = current_absences(home_id) if feature_on("absences") else []
    if home_id == 0:
        away = [x for x in away if x["a"]["home_id"] is None]
    log = db.execute("SELECT * FROM daily_logs WHERE date=? AND home_id IS ?", (d.isoformat(), log_home)).fetchone()
    # 同じ日・時間帯に記録がいくつかあるときは、いちばん新しいもの（番号が大きいもの）を表に出す
    recs = {r["resident_id"]: r for r in db.execute(
        "SELECT * FROM support_records WHERE date=? AND time_slot=? ORDER BY id", (d.isoformat(), slot))}
    if home_id == 0:
        day_all = db.execute("SELECT s.*, r.name AS rname FROM support_records s JOIN residents r ON r.id=s.resident_id "
                             "WHERE s.date=? AND r.home_id IS NULL ORDER BY r.kana, s.id", (d.isoformat(),)).fetchall()
    else:
        day_all = db.execute("SELECT s.*, r.name AS rname FROM support_records s JOIN residents r ON r.id=s.resident_id "
                             "WHERE s.date=? AND (? IS NULL OR r.home_id=?) ORDER BY r.kana, s.id",
                             (d.isoformat(), home_id, home_id)).fetchall()
    no_home = db.execute(f"SELECT COUNT(*) FROM residents WHERE home_id IS NULL AND ({ACTIVE_RES})").fetchone()[0]
    return render_template("journal.html", away=away, d=d, homes=homes, home_id=home_id, slot=slot, residents=residents, log=log,
                           recs=recs, day_all=day_all, prev=(d - timedelta(days=1)).isoformat(), no_home=no_home,
                           next=(d + timedelta(days=1)).isoformat(), MEAL=MEAL, MED=MED, MOOD=MOOD, TIME_SLOT=TIME_SLOT)


def _record_rows(rows):
    return [[r["date"], r["rname"], r["time_slot"] or "", r["temperature"] if r["temperature"] is not None else "",
             r["meal"] or "", r["medication"] or "", r["mood"] or "", r["content"] or "", r["staff"] or ""] for r in rows]


REC_HEAD = ["日付", "氏名", "時間帯", "体温", "食事", "服薬", "様子", "記録内容", "記録者"]
REC_WIDTH = [12, 14, 8, 7, 8, 8, 10, 60, 12]


@bp.route("/journal/export.xlsx")
def journal_export():
    db = get_db()
    first, last = parse_ym(request.args.get("ym"))
    home_id = request.args.get("home_id", type=db_int)
    home = db.execute("SELECT * FROM homes WHERE id=?", (home_id,)).fetchone() if home_id else None
    hname = home["name"] if home else "全住居"
    wb = Workbook()
    ws = wb.active
    ws.title = "業務日誌"
    logs = db.execute("SELECT * FROM daily_logs WHERE date BETWEEN ? AND ? AND (? IS NULL OR home_id=?) ORDER BY date",
                      (first.isoformat(), last.isoformat(), home_id, home_id)).fetchall()
    excel.add_table(ws, f"業務日誌　{first:%Y年%m月}　{hname}",
                    ["日付", "日中勤務者", "夜間勤務者", "在籍者数", "外泊・入院者", "全体の様子", "行事・来訪・通院", "申し送り", "確認者"],
                    [[l["date"], l["day_staff"] or "", l["night_staff"] or "", l["residents_count"] or "", l["absent"] or "",
                      l["summary"] or "", l["events"] or "", l["handover"] or "", l["checker"] or ""] for l in logs],
                    subtitle=get_setting("office_name"), widths=[12, 14, 14, 8, 16, 45, 30, 40, 10])
    ws2 = wb.create_sheet("支援記録")
    recs = db.execute("SELECT s.*, r.name AS rname FROM support_records s JOIN residents r ON r.id=s.resident_id "
                      "WHERE s.date BETWEEN ? AND ? AND (? IS NULL OR r.home_id=?) ORDER BY s.date, r.kana, s.id",
                      (first.isoformat(), last.isoformat(), home_id, home_id)).fetchall()
    excel.add_table(ws2, f"支援記録　{first:%Y年%m月}　{hname}", REC_HEAD, _record_rows(recs),
                    subtitle=get_setting("office_name"), widths=REC_WIDTH)
    return excel.send_workbook(wb, f"業務日誌_{first:%Y%m}_{hname}")


@bp.route("/residents/<int:rid>/records.xlsx")
def resident_records_export(rid):
    db = get_db()
    res = db.execute("SELECT * FROM residents WHERE id=?", (rid,)).fetchone()
    if res is None:
        abort(404)
    first, last = parse_ym(request.args.get("ym"))
    recs = db.execute("SELECT s.*, ? AS rname FROM support_records s WHERE resident_id=? AND date BETWEEN ? AND ? "
                      "ORDER BY date, id", (res["name"], rid, first.isoformat(), last.isoformat())).fetchall()
    return excel.send_table(f"支援記録_{res['name']}_{first:%Y%m}", REC_HEAD, _record_rows(recs),
                            subtitle=f"{get_setting('office_name')}　{first:%Y年%m月}")


# ---------------------------------------------------------------- 加算 要件確認・概算
def _addon_counts(first, last):
    """利用者別加算の対象人数（当月に有効なもの）"""
    rows = get_db().execute(
        f"SELECT ra.addon_id, COUNT(DISTINCT ra.resident_id) FROM resident_addons ra JOIN residents r ON r.id=ra.resident_id "
        "WHERE (ra.start_on IS NULL OR ra.start_on <= ?) AND (ra.end_on IS NULL OR ra.end_on >= ?) GROUP BY ra.addon_id",
        (last.isoformat(), first.isoformat())).fetchall()
    return {r[0]: r[1] for r in rows}


@bp.route("/addons/check", methods=["GET", "POST"])
def addon_check():
    db = get_db()
    first, last = parse_ym(request.values.get("ym"))
    ym = first.strftime("%Y-%m")
    addons = db.execute("SELECT * FROM addons WHERE active=1 ORDER BY id").fetchall()
    if request.method == "POST":
        for a in addons:
            for i, item in enumerate(_lines(a["requirements"])):
                checked = 1 if request.form.get(f"a{a['id']}_{i}") else 0
                db.execute("INSERT OR REPLACE INTO addon_checks (addon_id, ym, item, checked, checked_by, checked_at)"
                           " VALUES (?,?,?,?,?,?)", (a["id"], ym, item, checked, g.user["username"], now()))
        db.commit()
        flash("要件確認を保存しました。", "ok")
        return redirect(url_for("views.addon_check", ym=ym))
    checks = {(r["addon_id"], r["item"]): r for r in db.execute("SELECT * FROM addon_checks WHERE ym=?", (ym,))}
    counts = _addon_counts(first, last)
    n_res = db.execute(f"SELECT COUNT(*) FROM residents WHERE {ACTIVE_RES}").fetchone()[0]
    price = setting_number("unit_price", 10)
    days = (last - first).days + 1
    items, total_units = [], 0
    for a in addons:
        individual = (a["kind"] or "").startswith("個別")
        count = counts.get(a["id"], 0) if individual else n_res
        mult = days if a["unit_type"] == "日" else 1
        units = (a["units"] or 0) * count * mult
        if a["unit_type"] != "回":
            total_units += units
        reqs = [(i, item, checks.get((a["id"], item))) for i, item in enumerate(_lines(a["requirements"]))]
        items.append({"a": a, "count": count, "units": units, "reqs": reqs, "individual": individual,
                      "ok": all(c and c["checked"] for _, _, c in reqs)})
    return render_template("addon_check.html", items=items, ym=ym, first=first, total_units=total_units,
                           price=price, days=days, n_res=n_res)


def _lines(text):
    return [l.strip() for l in (text or "").splitlines() if l.strip()]


# ---------------------------------------------------------------- 処遇改善
@bp.route("/shogu/")
@admin_required
def shogu_index():
    plans = get_db().execute("SELECT * FROM shogu_plans ORDER BY fiscal_year DESC").fetchall()
    return render_template("shogu_index.html", plans=[(p, _shogu_summary(p)) for p in plans])


def _shogu_summary(plan):
    db = get_db()
    allocs = db.execute("SELECT a.*, s.name AS sname, s.job, s.employment FROM shogu_allocations a "
                        "LEFT JOIN staff s ON s.id=a.staff_id WHERE a.plan_id=? ORDER BY s.kana", (plan["id"],)).fetchall()
    addon = round((plan["revenue"] or 0) * (plan["rate"] or 0) / 100)
    monthly_total = sum((a["monthly"] or 0) * 12 for a in allocs)
    lump_total = sum(a["annual_lump"] or 0 for a in allocs)
    total = monthly_total + lump_total
    return {
        "allocs": allocs, "addon": addon, "monthly_total": monthly_total, "lump_total": lump_total, "total": total,
        "diff": total - addon,
        "monthly_ratio": (monthly_total / addon * 100) if addon else None,
    }


def _shogu_requirements(plan):
    reqs = get_db().execute("SELECT * FROM shogu_requirements ORDER BY sort, id").fetchall()
    cat = plan["category"] or ""
    try:
        checked = set(json.loads(plan["checked"] or "[]"))
    except ValueError:
        checked = set()
    out = []
    for r in reqs:
        targets = [t.strip() for t in (r["applies_to"] or "").replace("、", ",").split(",") if t.strip()]
        if cat and targets and cat not in targets:
            continue
        out.append((r, r["id"] in checked))
    return out


@bp.route("/shogu/<int:pid>", methods=["GET", "POST"])
@admin_required
def shogu_plan(pid):
    db = get_db()
    plan = db.execute("SELECT * FROM shogu_plans WHERE id=?", (pid,)).fetchone()
    if plan is None:
        abort(404)
    if request.method == "POST":
        ids = [int(x) for x in request.form.getlist("req")]
        db.execute("UPDATE shogu_plans SET checked=?, updated_at=?, updated_by=? WHERE id=?",
                   (json.dumps(ids), now(), g.user["username"], pid))
        audit("update", "shogu_plans", pid)
        db.commit()
        flash("要件の達成状況を保存しました。", "ok")
        return redirect(url_for("views.shogu_plan", pid=pid))
    return render_template("shogu_plan.html", plan=plan, s=_shogu_summary(plan), reqs=_shogu_requirements(plan),
                           label=label_of("shogu_plans", plan))


@bp.route("/shogu/<int:pid>/export.xlsx")
@admin_required
def shogu_export(pid):
    plan = get_db().execute("SELECT * FROM shogu_plans WHERE id=?", (pid,)).fetchone()
    if plan is None:
        abort(404)
    s = _shogu_summary(plan)
    title = f"処遇改善 配分表 {label_of('shogu_plans', plan)}"
    wb = Workbook()
    ws = wb.active
    ws.title = "職員別配分"
    rows = [[a["sname"] or "", a["job"] or "", a["employment"] or "", a["method"] or "", a["monthly"] or 0,
             (a["monthly"] or 0) * 12, a["annual_lump"] or 0, (a["monthly"] or 0) * 12 + (a["annual_lump"] or 0),
             a["notes"] or ""] for a in s["allocs"]]
    rows.append(["合計", "", "", "", "", s["monthly_total"], s["lump_total"], s["total"], ""])
    r = excel.add_table(ws, title, ["氏名", "職種", "雇用形態", "支給方法", "月額", "月額×12", "賞与・一時金", "年間合計", "備考"],
                        rows, subtitle=get_setting("office_name"), widths=[14, 16, 10, 14, 10, 12, 12, 12, 30])
    summary = [("加算率（%）", plan["rate"] or 0), ("報酬総額見込", plan["revenue"] or 0), ("加算額見込", s["addon"]),
               ("賃金改善額合計", s["total"]), ("差額（改善額−加算額）", s["diff"]),
               ("うち月額改善の割合（%）", round(s["monthly_ratio"], 1) if s["monthly_ratio"] is not None else "")]
    for i, (k, v) in enumerate(summary, r + 2):
        ws.cell(row=i, column=1, value=k)
        c = ws.cell(row=i, column=2, value=v)
        c.number_format = "#,##0.0" if isinstance(v, float) and v < 1000 else "#,##0"
    ws2 = wb.create_sheet("要件チェック")
    excel.add_table(ws2, f"処遇改善 要件チェック {label_of('shogu_plans', plan)}", ["要件", "対象区分", "内容", "達成"],
                    [[q["name"], q["applies_to"] or "", q["description"] or "", "✓" if ok else ""]
                     for q, ok in _shogu_requirements(plan)], widths=[34, 12, 70, 8])
    return excel.send_workbook(wb, title)


# ---------------------------------------------------------------- キャリアパス
def _career_rows():
    db = get_db()
    today = date.today()
    fy = fiscal_year(today)
    fy_start, fy_end = date(fy, 4, 1).isoformat(), date(fy + 1, 3, 31).isoformat()
    grades = dict(ref_options("career_grades"))
    rows = []
    for s in db.execute("SELECT * FROM staff WHERE status IS NULL OR status != '退職' ORDER BY kana"):
        hours = db.execute("SELECT COALESCE(SUM(hours),0), COUNT(*) FROM trainings WHERE staff_id=? AND date BETWEEN ? AND ?",
                           (s["id"], fy_start, fy_end)).fetchone()
        ev = db.execute("SELECT * FROM evaluations WHERE staff_id=? ORDER BY date DESC LIMIT 1", (s["id"],)).fetchone()
        tenure = years_between(parse_date(s["hire_date"]), today)
        rows.append({
            "s": s, "grade": grades.get(s["grade_id"], "未設定"), "tenure": tenure,
            "total_exp": (tenure or 0) + int(s["experience_years"] or 0),
            "hours": hours[0], "count": hours[1], "ev": ev,
            "recommend": grades.get(ev["recommend_grade_id"]) if ev else None,
        })
    return rows, fy


@bp.route("/career")
@admin_required
def career():
    db = get_db()
    rows, fy = _career_rows()
    grades = db.execute("SELECT g.*, (SELECT COUNT(*) FROM staff s WHERE s.grade_id=g.id AND (s.status IS NULL OR s.status!='退職')) AS n "
                        "FROM career_grades g ORDER BY level").fetchall()
    plans = db.execute("SELECT * FROM training_plans WHERE fiscal_year=? ORDER BY month", (fy,)).fetchall()
    return render_template("career.html", rows=rows, grades=grades, fy=fy, plans=plans)


@bp.route("/career/export.xlsx")
@admin_required
def career_export():
    db = get_db()
    rows, fy = _career_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "キャリアパス表"
    excel.add_table(ws, "キャリアパス（等級・任用要件・賃金体系）",
                    ["等級", "名称", "職位", "職責・職務内容", "任用要件", "必須研修", "基本給 下限", "基本給 上限", "役職手当", "昇給・昇格の基準"],
                    [[g["level"], g["name"], g["position"] or "", g["duties"] or "", g["requirements"] or "", g["trainings"] or "",
                      g["salary_min"] or "", g["salary_max"] or "", g["allowance"] or "", g["raise_rule"] or ""]
                     for g in db.execute("SELECT * FROM career_grades ORDER BY level")],
                    subtitle=get_setting("office_name"), widths=[6, 10, 16, 36, 30, 30, 12, 12, 10, 30])
    ws2 = wb.create_sheet("職員別状況")
    excel.add_table(ws2, f"職員別 等級・研修受講状況（{fy}年度）",
                    ["氏名", "職種", "雇用形態", "等級", "勤続年数", "通算経験年数", "保有資格", "研修受講数", "研修時間", "直近評価", "推薦等級"],
                    [[r["s"]["name"], r["s"]["job"] or "", r["s"]["employment"] or "", r["grade"],
                      r["tenure"] if r["tenure"] is not None else "", r["total_exp"], r["s"]["qualifications"] or "",
                      r["count"], r["hours"], (r["ev"]["score"] or "") if r["ev"] else "", r["recommend"] or ""] for r in rows])
    ws3 = wb.create_sheet("研修計画")
    excel.add_table(ws3, f"{fy}年度 研修計画", ["実施月", "研修名", "対象", "ねらい", "状況"],
                    [[p["month"] or "", p["title"], p["target"] or "", p["goal"] or "", p["status"] or ""]
                     for p in db.execute("SELECT * FROM training_plans WHERE fiscal_year=? ORDER BY month", (fy,))],
                    widths=[8, 30, 20, 50, 10])
    return excel.send_workbook(wb, "キャリアパス")


# ---------------------------------------------------------------- 帳票・設定・バックアップ
@bp.route("/reports")
def reports():
    db = get_db()
    return render_template("reports.html", homes=db.execute("SELECT * FROM homes ORDER BY name").fetchall(),
                           residents=db.execute(f"SELECT * FROM residents WHERE {ACTIVE_RES} ORDER BY kana").fetchall(),
                           plans=db.execute("SELECT * FROM shogu_plans ORDER BY fiscal_year DESC").fetchall(),
                           ym=date.today().strftime("%Y-%m"), label_of=label_of)


SETTINGS = [
    ("office_name", "事業所名"),
    ("office_no", "事業所番号"),
    ("unit_price", "1単位あたりの単価（円・地域区分に応じて）"),
    ("cert_alert_days", "受給者証期限アラート（何日前から）"),
    ("plan_alert_days", "個別支援計画・モニタリングのアラート（何日前から）"),
    ("office_address", "事業所の住所（請求書に印字）"),
    ("office_tel", "事業所の電話番号（請求書に印字）"),
    ("bank_info", "振込先（例：〇〇銀行 △△支店 普通 1234567 カ）〇〇）"),
    ("invoice_due_day", "利用料の支払期限（翌月の何日）"),
    ("full_time_hours", "常勤の勤務時間（月・時間）※常勤換算に使います"),
    ("corp_name", "法人の名前（運営規程・書類に印字）"),
    ("corp_rep", "法人の代表者（役職・氏名）"),
    ("service_area", "通常の事業の実施地域（例：〇〇市〇〇区）"),
    ("target_disability", "主たる対象とする障害の種類（例：知的障害者・精神障害者）"),
    ("complaint_staff", "苦情受付担当者（役職・氏名）"),
    ("complaint_manager", "苦情解決責任者（役職・氏名）"),
    ("third_party", "第三者委員（氏名・連絡先）"),
    ("cooperating_hospital", "協力医療機関（名前・診療科）"),
    ("session_timeout_min", "自動ログアウトまでの時間（分）※操作がないとき", ["10", "15", "30", "60", "120"]),
    ("pin_timeout_min", "PINでログインしたときの自動ログアウト（分）", ["5", "10", "15", "30", "60"]),
    ("staff_can_export", "職員のExcel出力・バックアップ", [("0", "許可しない（管理者だけ）"), ("1", "許可する")]),
]
SETTING_DEFAULTS = {"session_timeout_min": "30", "staff_can_export": "0", "pin_timeout_min": "15"}


@bp.route("/settings", methods=["GET", "POST"])
@admin_required
def settings():
    if request.method == "POST":
        bad = []
        for k, label, *opts in SETTINGS:
            v = request.form.get(k, "").strip()
            if opts and v not in [o if isinstance(o, str) else o[0] for o in opts[0]]:
                if v:  # 送られてこなかったときは、前の値のまま
                    bad.append(f"「{label}」はえらんだ値が正しくありません")
                continue
            if k in NUM_SETTINGS:
                v, err = check_setting(k, v)
                if err:  # まちがった値は保存せず、前の値のままにする
                    bad.append(f"「{label}」は{err}")
                    continue
            set_setting(k, v[:500])
        from .auth import log_event

        log_event("settings")
        get_db().commit()
        if bad:
            flash("次の項目は前の値のままにしました：" + "／".join(bad) + "。", "error")
        else:
            flash("設定を保存しました。", "ok")
        return redirect(url_for("views.settings"))
    items = []
    for k, label, *opts in SETTINGS:
        choices = [(o, o) if isinstance(o, str) else o for o in opts[0]] if opts else None
        items.append((k, label, get_setting(k, SETTING_DEFAULTS.get(k, "")), choices, NUM_SETTINGS.get(k)))
    return render_template("settings.html", items=items)


@bp.route("/backup")
@admin_required
def backup():
    src = sqlite3.connect(current_app.config["DATABASE"])
    mem = sqlite3.connect(":memory:")
    src.backup(mem)
    src.close()
    data = mem.serialize()
    mem.close()
    name = f"ghms_backup_{datetime.now():%Y%m%d_%H%M}.sqlite3"
    return send_file(io.BytesIO(data), mimetype="application/octet-stream", as_attachment=True, download_name=name)


@bp.route("/audit")
@admin_required
def audit_log():
    where, params = [], []
    user, action = request.args.get("user", ""), request.args.get("action", "")
    if user:
        where.append("username = ?")
        params.append(user)
    if action:
        where.append("action = ?")
        params.append(action)
    if request.args.get("from"):
        where.append("at >= ?")
        params.append(request.args["from"])
    if request.args.get("to"):
        where.append("at <= ?")
        params.append(request.args["to"] + " 23:59:59")
    sql = "SELECT * FROM audit_log" + (" WHERE " + " AND ".join(where) if where else "") + " ORDER BY id DESC LIMIT 500"
    db = get_db()
    rows = db.execute(sql, params).fetchall()
    users = [r[0] for r in db.execute("SELECT DISTINCT username FROM audit_log WHERE username != '' ORDER BY username")]
    return render_template("audit.html", rows=rows, users=users, args=request.args, ACTIONS=AUDIT_ACTIONS)


AUDIT_ACTIONS = {
    "login": "ログイン", "logout": "ログアウト", "login_failed": "ログイン失敗", "view": "閲覧", "export": "Excel出力・ダウンロード",
    "create": "登録", "update": "更新", "delete": "削除", "password_change": "パスワード変更", "settings": "設定の変更",
    "user_add": "ユーザー追加", "user_password": "パスワード再設定", "user_role": "権限の変更", "user_disable": "ユーザー停止",
    "user_enable": "ユーザー再開", "user_unlock": "ロック解除", "reauth": "管理者画面の本人確認",
    "pin_set": "PINの設定", "pin_clear": "PINを消した", "user_pin_clear": "PINを消した（管理者）",
    "device_add": "PIN端末の登録", "device_remove": "PIN端末の解除",
    "update_check": "更新の確認", "update_install": "更新の開始", "update_failed": "更新の失敗",
}
