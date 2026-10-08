"""提出・保管する書類をかんたんに作る。

- 献立表（住居ごと・1週間。印刷とExcel）
- サービス提供実績記録票（入居者ごと・月ごと。全員分をまとめて印刷）
- 指定更新の書類（提出書類のチェックリストと、データから作る書類）
  従業者の一覧／勤務体制（常勤換算）／管理者・サービス管理責任者の経歴書／運営規程のたたき台／
  利用者の状況／委員会・研修の実施状況

様式は自治体によって違うため、ここで作るのは「データを正しくそろえた下書き・控え」。
指定の様式がある場合は、ここの内容をもとに書き写してください。
"""

import calendar
from collections import Counter
from datetime import date, timedelta

from flask import Blueprint, abort, flash, g, redirect, render_template, request, url_for
from openpyxl import Workbook

from . import excel
from .auth import admin_required, log_event
from .billing import attendance_map, in_residence, month_days, no_home_count, pick_home, residents_in_month, sync_open
from .db import get_db, get_setting, now, set_setting
from .views import parse_date, parse_ym

bp = Blueprint("docs", __name__, url_prefix="/docs")

MEALS = ["朝食", "昼食", "夕食", "おやつ"]
WEEK = "月火水木金土日"


def office():
    keys = ["office_name", "office_no", "office_address", "office_tel", "corp_name", "corp_rep", "service_area",
            "target_disability", "complaint_staff", "complaint_manager", "third_party", "cooperating_hospital"]
    return {k: get_setting(k) for k in keys}


def _week_start(d):
    return d - timedelta(days=d.weekday())


# ---------------------------------------------------------------- 献立表
@bp.route("/menus", methods=["GET", "POST"])
def menus():
    db = get_db()
    homes = db.execute("SELECT * FROM homes ORDER BY name").fetchall()
    home_id = request.values.get("home_id", type=int) or (homes[0]["id"] if homes else None)
    start = _week_start(parse_date(request.values.get("week"), date.today()))
    days = [start + timedelta(days=i) for i in range(7)]
    if request.method == "POST":
        if request.form.get("action") == "copy_prev":
            prev = start - timedelta(days=7)
            rows = db.execute("SELECT * FROM menus WHERE home_id IS ? AND date BETWEEN ? AND ?",
                              (home_id, prev.isoformat(), (prev + timedelta(days=6)).isoformat())).fetchall()
            for r in rows:
                d = parse_date(r["date"]) + timedelta(days=7)
                db.execute("INSERT OR REPLACE INTO menus (home_id, date, meal, text, updated_by, updated_at) VALUES (?,?,?,?,?,?)",
                           (home_id, d.isoformat(), r["meal"], r["text"], g.user["username"], now()))
            flash(f"前の週の献立を写しました（{len(rows)}件）。必要なところを直して保存してください。", "ok")
        else:
            for d in days:
                for m in MEALS:
                    text = request.form.get(f"{d.isoformat()}_{m}", "").strip()
                    if text:
                        db.execute("INSERT OR REPLACE INTO menus (home_id, date, meal, text, updated_by, updated_at) VALUES (?,?,?,?,?,?)",
                                   (home_id, d.isoformat(), m, text, g.user["username"], now()))
                    else:
                        db.execute("DELETE FROM menus WHERE home_id IS ? AND date=? AND meal=?", (home_id, d.isoformat(), m))
            log_event("update", "menus", None, f"{start.isoformat()}の週")
            flash("献立を保存しました。", "ok")
        db.commit()
        return redirect(url_for("docs.menus", home_id=home_id or "", week=start.isoformat()))
    cells = {(r["date"], r["meal"]): r["text"] for r in db.execute(
        "SELECT * FROM menus WHERE home_id IS ? AND date BETWEEN ? AND ?", (home_id, days[0].isoformat(), days[-1].isoformat()))}
    allergies = db.execute("SELECT name, allergy FROM residents WHERE (status IS NULL OR status != '退居') AND home_id IS ? "
                           "AND allergy IS NOT NULL AND allergy != '' ORDER BY kana", (home_id,)).fetchall()
    home = next((h for h in homes if h["id"] == home_id), None)
    tpl = "menus_print.html" if request.args.get("print") else "menus.html"
    return render_template(tpl, homes=homes, home=home, home_id=home_id, days=days, start=start, cells=cells, MEALS=MEALS,
                           WEEK=WEEK, allergies=allergies, prev=(start - timedelta(days=7)).isoformat(),
                           next=(start + timedelta(days=7)).isoformat(), office=office())


@bp.route("/menus.xlsx")
def menus_export():
    db = get_db()
    first, last = parse_ym(request.args.get("ym"))
    home_id = request.args.get("home_id", type=int)
    home = db.execute("SELECT name FROM homes WHERE id=?", (home_id,)).fetchone() if home_id else None
    cells = {(r["date"], r["meal"]): r["text"] for r in db.execute(
        "SELECT * FROM menus WHERE home_id IS ? AND date BETWEEN ? AND ?", (home_id, first.isoformat(), last.isoformat()))}
    rows = [[f"{d.month}/{d.day}", WEEK[d.weekday()]] + [cells.get((d.isoformat(), m), "") for m in MEALS]
            for d in month_days(first, last)]
    return excel.send_table(f"献立表_{first:%Y%m}_{home['name'] if home else ''}", ["日付", "曜日"] + MEALS, rows,
                            subtitle=f"{get_setting('office_name')}　{home['name'] if home else ''}　{first:%Y年%m月}")


# ---------------------------------------------------------------- サービス提供実績記録票
# 列（日中支援・夜間支援など）は事業所で増やせる。
#   builtin="day" … 実績の「日」の日に自動（日中支援加算）
#   auto="stay" … 在居の日は自動でつく（夜間支援等体制加算など）。日ごとに外せる
#   auto="home" / "hosp" … 実績が「帰」／「入」の日は自動でつく
#   auto=""     … 日ごとに手でつける
# 記録票に出す記号は列ごとに決める（様式18-1の記載例では「1」）
BUILTIN_COLUMNS = [("day", "日中支援加算", "日")]
# 様式18-1（共同生活援助サービス提供実績記録票）の並び。（名前, つけ方, builtin, 合計の単位）
RECORD_DEFAULTS = [
    ("住居外利用", "", None, "日"),
    ("退居後支援", "", None, "回"),
    ("夜間支援等体制加算", "stay", None, "回"),
    ("入院時支援特別加算", "", None, "回"),
    ("帰宅時支援加算", "", None, "回"),
    ("日中支援加算", "", "day", "回"),
    ("医療連携体制加算", "", None, "回"),
    ("自立生活支援加算（Ⅰ）", "", None, "回"),
    ("自立生活支援加算（Ⅱ）", "", None, "回"),
    ("集中的支援加算", "", None, "回"),
]
RECORD_LAYOUT = "18-1"
AUTO_KINDS = {"": "日ごとに手でつける", "stay": "在居の日は自動でつく（日ごとに外せる）",
              "home": "実績が「帰」の日は自動でつく", "hosp": "実績が「入」の日は自動でつく"}
AUTO_CODES = {"stay": ("○", "日"), "home": ("帰",), "hosp": ("入",)}
DEFAULT_MARK = "1"
STAY_CODES = ("○", "日")
OUT_CODES = ("外", "帰")


def mark_text(col):
    return (col["mark"] or "").strip() or DEFAULT_MARK


def seed_record_columns(con):
    """様式18-1の列を用意する。前の並び（日中支援・夜間支援など4列）で、まだ何も入れていなければ入れ替える"""
    layout = con.execute("SELECT value FROM settings WHERE key='record_layout'").fetchone()
    if layout and layout[0] == RECORD_LAYOUT:
        return
    has_marks = con.execute("SELECT COUNT(*) FROM record_marks").fetchone()[0]
    if not has_marks:
        con.execute("DELETE FROM record_columns")
        for i, (label, auto, builtin, unit) in enumerate(RECORD_DEFAULTS):
            con.execute("INSERT INTO record_columns (label, sort, active, auto, builtin, mark, unit) VALUES (?,?,1,?,?,?,?)",
                        (label, (i + 1) * 10, auto, builtin, DEFAULT_MARK, unit))
    con.execute("INSERT OR REPLACE INTO settings (key, value) VALUES ('record_layout', ?)", (RECORD_LAYOUT,))
    con.commit()


def record_columns(active_only=True):
    sql = "SELECT * FROM record_columns" + (" WHERE active=1" if active_only else "") + " ORDER BY sort, id"
    return get_db().execute(sql).fetchall()


def _auto_mark(col, code):
    """手で直していないときに印がつくか"""
    if col["builtin"]:
        return code == dict((k, c) for k, _, c in BUILTIN_COLUMNS).get(col["builtin"])
    return code in AUTO_CODES.get(col["auto"] or "", ())


def marks_map(first, last):
    rows = get_db().execute("SELECT * FROM record_marks WHERE date BETWEEN ? AND ?", (first.isoformat(), last.isoformat()))
    return {(m["resident_id"], m["date"], m["col_id"]): m["value"] for m in rows}


def mark_of(col, code, stored):
    if col["builtin"] or stored is None:
        return _auto_mark(col, code)
    return stored == "○"


def _place(code):
    if code in OUT_CODES:
        return "外泊"
    if code == "入":
        return "入院"
    return "home" if code in STAY_CODES else None


def status_label(prev, code, home):
    """サービス提供の状況（記載例：「〇〇ホーム→外泊」「外泊」「外泊戻り」。ふつうに支援した日は空欄）"""
    cur, before = _place(code), _place(prev)
    if cur is None:
        return ""
    if cur == "home":
        return f"{before}戻り" if before in ("外泊", "入院") else ""
    if before == cur:
        return cur
    return f"{before if before in ('外泊', '入院') else (home or '住居')}→{cur}"


def record_sheet_data(first, last, rid=None):
    sync_open()
    residents = [r for r in residents_in_month(first, last) if rid is None or r["id"] == rid]
    amap = attendance_map(first - timedelta(days=1), last)
    mmap = marks_map(first, last)
    cols = record_columns()
    sheets = []
    for r in residents:
        rows = []
        before = first - timedelta(days=1)
        prev = amap.get((r["id"], before.isoformat()), "") if in_residence(r, before) else ""
        for d in month_days(first, last):
            # 入居前・退居後の日は、実績も印もつけない
            inres = in_residence(r, d)
            code = amap.get((r["id"], d.isoformat()), "") if inres else ""
            marks = [mark_text(c) if inres and mark_of(c, code, mmap.get((r["id"], d.isoformat(), c["id"]))) else ""
                     for c in cols]
            rows.append({"d": d, "w": WEEK[d.weekday()], "code": code, "label": status_label(prev, code, r["hname"]),
                         "marks": marks})
            prev = code
        count = Counter(x["code"] for x in rows)
        rec = str(r["recipient_no"] or "").strip()
        sheets.append({"r": r, "rows": rows, "stay": count["○"] + count["日"], "home": count["帰"], "hosp": count["入"],
                       "out": count["外"], "empty": sum(1 for x in rows if not x["code"] and in_residence(r, x["d"])),
                       "rec_digits": list(rec.ljust(10))[:10] if len(rec) <= 10 else [rec],
                       "totals": [sum(1 for x in rows if x["marks"][i]) for i in range(len(cols))]})
    return sheets, cols


@bp.route("/record-sheets")
@admin_required
def record_sheets():
    first, last = parse_ym(request.args.get("ym"))
    rid = request.args.get("rid", type=int)
    sheets, cols = record_sheet_data(first, last, rid)
    log_event("view", "record_sheets", rid, f"{first:%Y-%m} 実績記録票")
    get_db().commit()
    return render_template("record_sheets.html", sheets=sheets, cols=cols, first=first, office=office(), ym=first.strftime("%Y-%m"),
                           residents=residents_in_month(first, last), rid=rid)


@bp.route("/record-sheets.xlsx")
@admin_required
def record_sheets_export():
    first, last = parse_ym(request.args.get("ym"))
    o = office()
    wb = Workbook()
    wb.remove(wb.active)
    sheets, cols = record_sheet_data(first, last, request.args.get("rid", type=int))
    for sh in sheets:
        r = sh["r"]
        ws = wb.create_sheet(excel.safe_sheet_title(r["name"]))
        rows = [[f"{x['d'].day}", x["w"], x["label"]] + x["marks"] + ["", ""] for x in sh["rows"]]
        rows.append(["合計", "", ""] + [f"{t}{c['unit'] or '回'}" for t, c in zip(sh["totals"], cols)] + ["", ""])
        excel.add_table(ws, f"サービス提供実績記録票（共同生活援助）　{first:%Y年%m月}分",
                        ["日付", "曜日", "サービス提供の状況"] + [c["label"] for c in cols] + ["利用者確認欄", "備考"], rows,
                        subtitle=f"事業所：{o['office_name']}（{o['office_no']}）　受給者証番号：{r['recipient_no'] or ''}　"
                                 f"支給決定障害者等氏名：{r['name']}　障害支援区分：{r['support_level'] or ''}",
                        widths=[5, 5, 18] + [9] * len(cols) + [9, 14])
    if not wb.sheetnames:
        wb.create_sheet("なし")
    return excel.send_workbook(wb, f"サービス提供実績記録票_{first:%Y%m}")


@bp.route("/record-columns", methods=["GET", "POST"])
@admin_required
def record_columns_settings():
    """実績記録票の項目を増やす・名前を変える・使わない・並べ替え"""
    db = get_db()
    if request.method == "POST":
        for c in record_columns(active_only=False):
            cid = c["id"]
            if request.form.get(f"delete::{cid}") and not c["builtin"]:
                db.execute("DELETE FROM record_columns WHERE id=?", (cid,))
                db.execute("DELETE FROM record_marks WHERE col_id=?", (cid,))
                continue
            label = (request.form.get(f"label::{cid}") or "").strip()[:20] or c["label"]
            auto = c["auto"] if c["builtin"] else (request.form.get(f"auto::{cid}") or "")
            mark = (request.form.get(f"mark::{cid}") or "").strip()[:3] or DEFAULT_MARK
            unit = "日" if request.form.get(f"unit::{cid}") == "日" else "回"
            db.execute("UPDATE record_columns SET label=?, sort=?, active=?, auto=?, mark=?, unit=? WHERE id=?",
                       (label, request.form.get(f"sort::{cid}", type=int) or 0, 1 if request.form.get(f"active::{cid}") else 0,
                        auto if auto in AUTO_KINDS else "", mark, unit, cid))
        new = (request.form.get("new") or "").strip()[:20]
        if new:
            mx = db.execute("SELECT COALESCE(MAX(sort),0) FROM record_columns").fetchone()[0]
            auto = request.form.get("new_auto") or ""
            db.execute("INSERT INTO record_columns (label, sort, active, auto, builtin, mark, unit) VALUES (?,?,1,?,NULL,?,?)",
                       (new, mx + 10, auto if auto in AUTO_KINDS else "", (request.form.get("new_mark") or "").strip()[:3] or DEFAULT_MARK,
                        "日" if request.form.get("new_unit") == "日" else "回"))
        db.commit()
        flash("実績記録票の項目を保存しました。", "ok")
        return redirect(url_for("docs.record_columns_settings"))
    used = {r[0]: r[1] for r in db.execute("SELECT col_id, COUNT(*) FROM record_marks GROUP BY col_id")}
    return render_template("record_columns.html", cols=record_columns(active_only=False), used=used, AUTO_KINDS=AUTO_KINDS,
                           DEFAULT_MARK=DEFAULT_MARK,
                           BUILTIN={k: c for k, _, c in BUILTIN_COLUMNS})


@bp.route("/record-marks", methods=["GET", "POST"])
def record_marks():
    """夜間支援など、増やした項目を月の表でチェックする（職員も入力できる）"""
    db = get_db()
    first, last = parse_ym(request.values.get("ym"))
    ym = first.strftime("%Y-%m")
    homes = db.execute("SELECT * FROM homes ORDER BY name").fetchall()
    home_id = pick_home(homes)
    cols = [c for c in record_columns() if not c["builtin"]]
    col = next((c for c in cols if c["id"] == request.values.get("col", type=int)), cols[0] if cols else None)
    residents = residents_in_month(first, last, home_id)
    days = month_days(first, last)
    if col is None:
        return render_template("record_marks.html", col=None, cols=cols, ym=ym, first=first, homes=homes, home_id=home_id,
                               no_home=no_home_count())
    sync_open()
    amap = attendance_map(first, last)
    if request.method == "POST":
        for r in residents:
            for d in days:
                code = amap.get((r["id"], d.isoformat()), "")
                checked = bool(request.form.get(f"m{r['id']}_{d.day}"))
                if not in_residence(r, d) or checked == _auto_mark(col, code):  # 自動と同じなら記録しない（実績を直したときに追いかける）
                    db.execute("DELETE FROM record_marks WHERE resident_id=? AND date=? AND col_id=?", (r["id"], d.isoformat(), col["id"]))
                else:
                    db.execute("INSERT OR REPLACE INTO record_marks (resident_id, date, col_id, value, updated_by, updated_at)"
                               " VALUES (?,?,?,?,?,?)", (r["id"], d.isoformat(), col["id"], "○" if checked else "-",
                                                         g.user["username"], now()))
        db.commit()
        flash(f"{first:%Y年%m月}の「{col['label']}」を保存しました。", "ok")
        return redirect(url_for("docs.record_marks", ym=ym, home_id=home_id, col=col["id"]))
    mmap = marks_map(first, last)
    rows = []
    for r in residents:
        cells = []
        for d in days:
            inres = in_residence(r, d)
            code = amap.get((r["id"], d.isoformat()), "") if inres else ""
            cells.append({"d": d, "code": code, "on": inres and mark_of(col, code, mmap.get((r["id"], d.isoformat(), col["id"]))),
                          "stay": code in STAY_CODES, "inres": inres})
        rows.append({"r": r, "cells": cells, "total": sum(1 for c in cells if c["on"])})
    return render_template("record_marks.html", col=col, cols=cols, ym=ym, first=first, homes=homes, home_id=home_id,
                           rows=rows, days=days, WEEK=WEEK, AUTO_KINDS=AUTO_KINDS, mark=mark_text(col), no_home=no_home_count())


# ---------------------------------------------------------------- 指定更新の書類
RENEWAL_DEFAULTS = [
    # (書類, システムで作れる画面のendpoint または None)
    ("指定更新申請書", None),
    ("付表（共同生活援助事業所の指定に係る記載事項）", None),
    ("従業者の勤務の体制及び勤務形態一覧表", "shift.index"),
    ("従業者の一覧（職種・資格・常勤換算）", "docs.staff_list"),
    ("管理者の経歴書", "docs.resumes"),
    ("サービス管理責任者の経歴書・研修修了証の写し", "docs.resumes"),
    ("資格証の写し（介護福祉士・社会福祉士など）", None),
    ("運営規程", "docs.rules"),
    ("利用者の状況（定員・現員・障害支援区分）", "docs.residents_status"),
    ("委員会・研修の実施状況（虐待防止・身体拘束・感染症・BCP）", "docs.committee"),
    ("事業所の平面図・写真", None),
    ("設備・備品等一覧表", None),
    ("協力医療機関との契約書の写し", None),
    ("利用者からの苦情を解決するための措置の概要", "docs.rules"),
    ("誓約書・役員等名簿", None),
    ("業務継続計画（BCP）", None),
]


def seed_renewal(con):
    if con.execute("SELECT COUNT(*) FROM renewal_items").fetchone()[0]:
        return
    for i, (name, ep) in enumerate(RENEWAL_DEFAULTS):
        con.execute("INSERT INTO renewal_items (sort, name, endpoint, done, note) VALUES (?,?,?,0,'')", ((i + 1) * 10, name, ep))
    con.commit()


@bp.route("/renewal", methods=["GET", "POST"])
@admin_required
def renewal():
    db = get_db()
    seed_renewal(db)
    if request.method == "POST":
        for r in db.execute("SELECT * FROM renewal_items").fetchall():
            db.execute("UPDATE renewal_items SET done=?, note=? WHERE id=?",
                       (1 if request.form.get(f"done_{r['id']}") else 0, request.form.get(f"note_{r['id']}", "")[:200], r["id"]))
        new = request.form.get("new", "").strip()
        if new:
            mx = db.execute("SELECT COALESCE(MAX(sort), 0) FROM renewal_items").fetchone()[0]
            db.execute("INSERT INTO renewal_items (sort, name, endpoint, done, note) VALUES (?,?,NULL,0,'')", (mx + 10, new[:80]))
        set_setting("renewal_due", request.form.get("due", "").strip())
        log_event("settings", detail="指定更新の書類チェックリスト")
        db.commit()
        flash("保存しました。", "ok")
        return redirect(url_for("docs.renewal"))
    items = db.execute("SELECT * FROM renewal_items ORDER BY sort, id").fetchall()
    return render_template("renewal.html", items=items, due=get_setting("renewal_due", ""),
                           done=sum(1 for i in items if i["done"]), office=office())


def _staff_rows():
    """在籍職員と、今月の勤務表からの常勤換算"""
    from .shift import roster

    first = date.today().replace(day=1)
    last = date(first.year, first.month, calendar.monthrange(first.year, first.month)[1])
    _, _, _, rows, by_job, _, base = roster(first, last, None)
    fte = {r["s"]["id"]: (r["fte"], r["hours"]) for r in rows}
    staff = get_db().execute("SELECT s.*, h.name AS hname FROM staff s LEFT JOIN homes h ON h.id=s.home_id "
                             "WHERE s.status IS NULL OR s.status != '退職' ORDER BY s.job, s.kana").fetchall()
    return staff, fte, by_job, first


@bp.route("/staff-list")
@admin_required
def staff_list():
    staff, fte, by_job, first = _staff_rows()
    return render_template("doc_staff_list.html", staff=staff, fte=fte, by_job=by_job, first=first, office=office(),
                           today=date.today())


@bp.route("/resumes")
@admin_required
def resumes():
    db = get_db()
    sid = request.args.get("sid", type=int)
    if sid:
        people = db.execute("SELECT * FROM staff WHERE id=?", (sid,)).fetchall()
    else:
        people = db.execute("SELECT * FROM staff WHERE (status IS NULL OR status != '退職') AND job IN ('管理者', 'サービス管理責任者') "
                            "ORDER BY job, kana").fetchall()
    log_event("view", "staff", sid, "経歴書")
    db.commit()
    all_staff = db.execute("SELECT id, name, job FROM staff WHERE status IS NULL OR status != '退職' ORDER BY kana").fetchall()
    return render_template("doc_resumes.html", people=people, office=office(), today=date.today(), all_staff=all_staff, sid=sid)


@bp.route("/rules")
@admin_required
def rules():
    db = get_db()
    homes = db.execute("SELECT * FROM homes ORDER BY name").fetchall()
    staff = db.execute("SELECT job, employment, COUNT(*) AS n FROM staff WHERE status IS NULL OR status != '退職' "
                       "GROUP BY job, employment").fetchall()
    jobs = {}
    for s in staff:
        jobs.setdefault(s["job"] or "その他", {})[s["employment"] or "—"] = s["n"]
    fees = db.execute("SELECT MIN(rent) AS rmin, MAX(rent) AS rmax, MIN(utility) AS umin, MAX(utility) AS umax, "
                      "MIN(daily_goods) AS gmin, MAX(daily_goods) AS gmax, MIN(food_amount) AS fmin, MAX(food_amount) AS fmax "
                      "FROM residents WHERE status IS NULL OR status != '退居'").fetchone()
    return render_template("doc_rules.html", o=office(), office=office(), homes=homes, capacity=int(sum(h["capacity"] or 0 for h in homes)),
                           jobs=jobs, fees=fees, types=get_setting("home_types", ""), today=date.today())


@bp.route("/residents-status")
@admin_required
def residents_status():
    db = get_db()
    homes = db.execute("SELECT * FROM homes ORDER BY name").fetchall()
    res = db.execute("SELECT * FROM residents WHERE status IS NULL OR status != '退居'").fetchall()
    levels = ["非該当", "区分1", "区分2", "区分3", "区分4", "区分5", "区分6"]
    table = []
    for h in homes:
        mine = [r for r in res if r["home_id"] == h["id"]]
        c = Counter(r["support_level"] or "未入力" for r in mine)
        table.append({"h": h, "n": len(mine), "levels": c})
    dis = Counter(r["disability_type"] or "未入力" for r in res)
    total_levels = Counter(r["support_level"] or "未入力" for r in res)
    return render_template("doc_residents_status.html", table=table, levels=levels, dis=dis, total=len(res),
                           total_levels=total_levels, office=office(), today=date.today(),
                           capacity=int(sum(h["capacity"] or 0 for h in homes)))


@bp.route("/committee")
@admin_required
def committee():
    db = get_db()
    since = (date.today() - timedelta(days=int(request.args.get("days", 730)))).isoformat()
    meetings = db.execute("SELECT * FROM meetings WHERE date >= ? ORDER BY kind, date", (since,)).fetchall()
    trainings = db.execute("SELECT t.*, s.name AS sname FROM trainings t LEFT JOIN staff s ON s.id=t.staff_id "
                           "WHERE t.date >= ? ORDER BY t.date", (since,)).fetchall()
    by_kind = {}
    for m in meetings:
        by_kind.setdefault(m["kind"] or "その他", []).append(m)
    return render_template("doc_committee.html", by_kind=by_kind, trainings=trainings, since=since, office=office(),
                           today=date.today())
