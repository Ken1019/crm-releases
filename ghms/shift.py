"""勤務表（勤務形態一覧表）と常勤換算。"""

from collections import defaultdict

from flask import Blueprint, flash, g, redirect, render_template, request, url_for

from . import excel
from .auth import admin_required
from .billing import month_days
from .db import get_db, get_setting, now
from .forms import db_int
from .views import parse_ym

bp = Blueprint("shift", __name__, url_prefix="/shift")

WEEK = "月火水木金土日"


def _staff(home_id):
    sql = "SELECT * FROM staff WHERE (status IS NULL OR status != '退職')"
    params = []
    if home_id:
        sql += " AND home_id=?"
        params.append(home_id)
    return get_db().execute(sql + " ORDER BY job, kana", params).fetchall()


def _types():
    return get_db().execute('SELECT * FROM shift_types ORDER BY sort, id').fetchall()


def full_time_hours():
    from .forms import setting_number

    return float(setting_number("full_time_hours", 160))


def roster(first, last, home_id):
    db = get_db()
    staff = _staff(home_id)
    types = _types()
    tmap = {t["code"]: t for t in types}
    days = month_days(first, last)
    smap = {(s["staff_id"], s["date"]): s["code"] for s in db.execute(
        "SELECT * FROM shifts WHERE date BETWEEN ? AND ?", (first.isoformat(), last.isoformat()))}
    fte_base = full_time_hours()
    rows, by_job = [], defaultdict(float)
    night_per_day = [0] * len(days)
    for s in staff:
        cells = [smap.get((s["id"], d.isoformat()), "") for d in days]
        hours = sum((tmap[c]["hours"] or 0) for c in cells if c in tmap)
        for i, c in enumerate(cells):
            if c in tmap and tmap[c]["night"]:
                night_per_day[i] += 1
        # 常勤の職員は1.0、非常勤は勤務時間 ÷ 常勤の勤務時間（上限1.0）
        fte = 1.0 if s["employment"] == "常勤" else min(round(hours / fte_base, 2), 1.0) if fte_base else 0
        by_job[s["job"] or "未設定"] += fte
        rows.append({"s": s, "cells": cells, "hours": hours, "fte": fte})
    return staff, types, days, rows, dict(by_job), night_per_day, fte_base


@bp.route("/", methods=["GET", "POST"])
@admin_required
def index():
    db = get_db()
    first, last = parse_ym(request.values.get("ym"))
    ym = first.strftime("%Y-%m")
    homes = db.execute("SELECT * FROM homes ORDER BY name").fetchall()
    home_id = request.values.get("home_id", type=db_int)
    staff, types, days, rows, by_job, night, base = roster(first, last, home_id)
    if request.method == "POST":
        valid = {t["code"] for t in types}
        for s in staff:
            db.execute("DELETE FROM shifts WHERE staff_id=? AND date BETWEEN ? AND ?", (s["id"], first.isoformat(), last.isoformat()))
            for d in days:
                code = request.form.get(f"s{s['id']}_{d.day}", "")
                if code in valid:
                    db.execute("INSERT INTO shifts (staff_id, date, code, updated_by, updated_at) VALUES (?,?,?,?,?)",
                               (s["id"], d.isoformat(), code, g.user["username"], now()))
        db.commit()
        flash(f"{first:%Y年%m月}の勤務表を保存しました。", "ok")
        return redirect(url_for("shift.index", ym=ym, home_id=home_id or ""))
    return render_template("shift.html", ym=ym, first=first, homes=homes, home_id=home_id, types=types, days=days,
                           rows=rows, by_job=by_job, night=night, base=base, WEEK=WEEK)


@bp.route("/export.xlsx")
@admin_required
def export():
    first, last = parse_ym(request.args.get("ym"))
    home_id = request.args.get("home_id", type=db_int)
    staff, types, days, rows, by_job, night, base = roster(first, last, home_id)
    home = get_db().execute("SELECT name FROM homes WHERE id=?", (home_id,)).fetchone() if home_id else None
    headers = ["職種", "雇用形態", "氏名"] + [f"{d.day}\n{WEEK[d.weekday()]}" for d in days] + ["勤務時間", "常勤換算"]
    data = [[r["s"]["job"] or "", r["s"]["employment"] or "", r["s"]["name"]] + r["cells"] + [r["hours"], r["fte"]] for r in rows]
    data.append(["", "", "夜間の勤務者数"] + night + ["", ""])
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "勤務表"
    end = excel.add_table(ws, f"勤務形態一覧表　{first:%Y年%m月}　{home['name'] if home else '全住居'}", headers, data,
                          subtitle=f"{get_setting('office_name')}　常勤の勤務時間：月{base:g}時間",
                          widths=[14, 8, 12] + [4] * len(days) + [9, 9])
    r = end + 2
    ws.cell(row=r, column=1, value="職種ごとの常勤換算")
    for job, v in by_job.items():
        r += 1
        ws.cell(row=r, column=1, value=job)
        ws.cell(row=r, column=3, value=round(v, 2))
    r += 2
    ws.cell(row=r, column=1, value="記号：" + "　".join(f"{t['code']}={t['name']}" + (f"({t['start']}〜{t['end']})" if t["start"] else "") for t in types))
    return excel.send_workbook(wb, f"勤務表_{first:%Y%m}")
