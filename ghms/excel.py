"""Excel（.xlsx）帳票の出力。"""

import io
from datetime import datetime
from urllib.parse import quote

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
FONT = "Meiryo UI"


def cell_value(field, value, maps):
    if value is None:
        return ""
    if field["type"] == "ref":
        return maps.get(field["name"], {}).get(value, "")
    if field["type"] == "multiref":
        m = maps.get(field["name"], {})
        return "、".join(m.get(int(i), "") for i in str(value).split(",") if i.strip().isdigit())
    if field["type"] == "check":
        return "✓" if value else ""
    return value


def add_table(ws, title, headers, rows, subtitle=None, start_row=1, widths=None):
    ws.cell(row=start_row, column=1, value=title).font = Font(name=FONT, size=14, bold=True)
    r = start_row + 1
    if subtitle:
        ws.cell(row=r, column=1, value=subtitle).font = Font(name=FONT, size=10, color="555555")
        r += 1
    r += 1
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_row = r
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(v, (int, float)) and abs(v) >= 1000:
                cell.number_format = "#,##0"
    for c, h in enumerate(headers, 1):
        if widths and c - 1 < len(widths):
            w = widths[c - 1]
        else:
            longest = max([len(str(h))] + [min(len(str(row[c - 1])), 40) for row in rows if c - 1 < len(row)])
            w = min(max(8, longest * 2 + 2), 60)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return r


def safe_sheet_title(s):
    for ch in '[]:*?/\\':
        s = s.replace(ch, "_")
    return s[:31] or "Sheet"


def send_workbook(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"{filename}_{datetime.now():%Y%m%d}.xlsx"
    resp = send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=name)
    resp.headers["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(name)}"
    return resp


def send_table(title, headers, rows, subtitle=None):
    wb = Workbook()
    ws = wb.active
    ws.title = safe_sheet_title(title)
    add_table(ws, title, headers, rows, subtitle or f"出力日時: {datetime.now():%Y-%m-%d %H:%M}")
    return send_workbook(wb, title)
