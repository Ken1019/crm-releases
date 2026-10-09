"""勤務の実績（出勤簿・タイムカード）と職員の体温（健康チェック）。

勤務の実績の作り方は2つ（設定の attend_mode。画面には出さない）：
- "shift"（初め）：管理者が勤務表を組み、その日に「出勤」（勤務表の時間で確定。時間は直せる）か
  「休み」「有給」を押す（勤務の確定 /work/confirm）。打刻の画面は出さない。体温は職員が自分の画面で入れる
- "punch"：事務所のPCで打刻する（/work/kiosk・/work/clock）。今は使っていないが、将来のために残してある。
  使うときは settings に attend_mode=punch を入れる（docs/販売に向けて.md）
打刻のとき：
- 職員は自分の画面で「出勤する」「退勤する」を押すだけ。出勤のときに体温と体調をたずねる
- 夜勤のように日をまたいだ勤務は、出勤した日の記録として扱う（退勤が出勤より前の時刻なら翌日）
- 管理者は職員ごと・月ごとに打刻を直せる（直した人と日時が残る）。出勤簿をExcelで出せる
"""

from datetime import date, datetime, timedelta

from flask import Blueprint, abort, flash, g, redirect, render_template, request, session, url_for

from . import excel
from .auth import admin_required, is_admin, log_event, require_admin
from .billing import month_days
from .db import get_db, get_setting, now
from .forms import clamp, db_int, finite_float, safe_int, setting_number
from .views import parse_date, parse_ym

bp = Blueprint("work", __name__, url_prefix="/work")

SYMPTOMS = ["せき", "のどの痛み", "鼻水", "だるさ", "頭痛", "下痢", "吐き気・おう吐", "味やにおいがわかりにくい"]
WEEK = "月火水木金土日"


def punch_mode():
    """勤務の実績の作り方。"shift"＝勤務表から確定する（初め）／"punch"＝事務所のPCで打刻する（今は画面に出さない）"""
    return "punch" if get_setting("attend_mode", "shift") == "punch" else "shift"


def setting_num(key, default):
    """数字の設定。まちがった値が保存されていても初期値を使い、決まった範囲の中におさめる（forms.NUM_SETTINGS）"""
    return float(setting_number(key, default))


def fever_line():
    return setting_num("pay_fever", 37.5)


def night_band():
    """深夜手当の時間帯（設定。初期は22時〜翌9時。法律の深夜割増は22時〜翌5時で、それより広くするのはよい）"""
    s = _min(get_setting("pay_night_start", "22:00"))
    e = _min(get_setting("pay_night_end", "09:00"))
    return (22 * 60 if s is None else s), (9 * 60 if e is None else e)


def night_windows():
    """出勤日の0時を0とした分で、深夜の時間帯を前の日〜翌々日の分まで並べる"""
    s, e = night_band()
    if e <= s:
        return [(k * 1440 + s, (k + 1) * 1440 + e) for k in (-1, 0, 1)]
    return [(k * 1440 + s, k * 1440 + e) for k in (0, 1, 2)]


def night_label():
    s, e = night_band()
    return f"{s // 60}時{f'{s % 60}分' if s % 60 else ''}〜{'翌' if e <= s else ''}{e // 60}時{f'{e % 60}分' if e % 60 else ''}"


def _min(hhmm):
    try:
        h, m = (hhmm or "").split(":")[:2]
        return int(h) * 60 + int(m)
    except ValueError:
        return None


def admin_view():
    """管理者の画面として開いてよいか（PINで入った管理者は、パスワードで確認するまでだめ）"""
    return is_admin() and session.get("via") != "pin"


def work_minutes(card):
    """1回の勤務の 実働・残業（この1回だけで8時間をこえた分）・深夜（設定の時間帯）・夜勤かどうか。
    夜勤は日をまたぐ勤務だけ（朝早い勤務は、深夜の時間帯に入っていても夜勤にしない）。
    月の残業は、1日（出勤した日）の合計で8時間をこえた分で数える（day_cards・month_summary）"""
    start, end = _min(card["clock_in"]), _min(card["clock_out"])
    if start is None or end is None:
        return None
    if end < start:
        end += 1440  # 同じ時刻（押してすぐ退勤）は0分。24時間の夜勤にはしない（1回の勤務は上限20時間）
    brk = card["break_min"] or 0
    total = max(0, end - start - brk)
    night = sum(max(0, min(end, b) - max(start, a)) for a, b in night_windows())
    return {"total": total, "over": max(0, total - 480), "night": max(0, night), "yakin": end > 1440}


def day_cards(cards):
    """打刻を日ごとにまとめて、残業を「1日の合計で8時間をこえた分」にする。
    日をまたぐ勤務は出勤した日の分。1日に何回か打刻したときは、8時間をこえたあとの打刻に残業をつける。
    [{"c": 打刻, "w": 実働など（over は1日で数えた分）}] を返す（退勤がない打刻は w=None）"""
    def key(c):
        m = _min(c["clock_in"])
        return (c["date"], -1 if m is None else m)

    out, done = [], {}
    for c in sorted(cards, key=key):
        w = work_minutes(c)
        if w is not None:
            before = done.get(c["date"], 0)
            after = before + w["total"]
            done[c["date"]] = after
            w = dict(w, over=max(0, after - 480) - max(0, before - 480))
        out.append({"c": c, "w": w})
    return out


def hm(minutes):
    if not minutes:
        return ""
    return f"{minutes // 60}:{minutes % 60:02d}"


def my_staff_id():
    """ログインしている人の職員ID（管理者が「ログインする人」の画面で職員の情報とつないだときだけ。
    名前が同じというだけではつながない。ほかの人の給与やタイムカードが見えてしまうため）"""
    return g.user["staff_id"] or None


def max_shift_minutes():
    """1回の勤務の上限。夜勤（例 16時〜翌10時）があるので設定で変えられる。これをこえた打刻は退勤を押せず、管理者が直す"""
    return int(setting_num("pay_max_shift_hours", 20) * 60)


def _started(card):
    try:
        return datetime.strptime(f"{card['date']} {card['clock_in']}", "%Y-%m-%d %H:%M")
    except (TypeError, ValueError):
        return None


def forgot_minutes():
    """出勤からこの時間をすぎて、また出勤を押したら、前の打刻は「退勤忘れ」にして新しく出勤する"""
    return int(setting_num("pay_forgot_hours", 12) * 60)


# あとに出勤の打刻がある（＝そのあと新しく出勤した）打刻。退勤がなければ「退勤忘れ」
LATER_CARD = ("EXISTS (SELECT 1 FROM timecards t2 WHERE t2.staff_id=timecards.staff_id AND t2.id != timecards.id"
              " AND (t2.date > timecards.date OR (t2.date = timecards.date AND COALESCE(t2.clock_in, '') > COALESCE(timecards.clock_in, ''))))")


def open_card(staff_id):
    """まだ退勤していない打刻。出勤から上限の時間をすぎたもの・あとに新しく出勤したものは「退勤忘れ」として扱い、ここでは返さない"""
    since = (date.today() - timedelta(days=2)).isoformat()
    for c in get_db().execute("SELECT * FROM timecards WHERE staff_id=? AND clock_out IS NULL AND date >= ? AND NOT " + LATER_CARD
                              + " ORDER BY date DESC, id DESC", (staff_id, since)):
        st = _started(c)
        if st and datetime.now() - st <= timedelta(minutes=max_shift_minutes()):
            return c
    return None


def is_stale(card):
    """出勤から「退勤忘れとみなす時間」をすぎた打刻か（夜勤明けの退勤か、退勤の押し忘れかを本人にえらんでもらう）"""
    st = _started(card) if card else None
    return bool(st and datetime.now() - st > timedelta(minutes=forgot_minutes()))


def forgotten_cards(staff_id=None, days=31):
    """退勤の打刻がないまま上限をすぎたもの・そのあと新しく出勤したもの（管理者が直す）"""
    since = (date.today() - timedelta(days=days)).isoformat()
    sql, args = f"SELECT *, {LATER_CARD} AS later FROM timecards WHERE clock_out IS NULL AND date >= ?", [since]
    if staff_id:
        sql += " AND staff_id=?"
        args.append(staff_id)
    out = []
    for c in get_db().execute(sql + " ORDER BY date", args):
        st = _started(c)
        if c["later"] or st is None or datetime.now() - st > timedelta(minutes=max_shift_minutes()):
            out.append(c)
    return out


def default_break(card, out_time):
    cin, cout = _min(card["clock_in"]), _min(out_time)
    if cin is None or cout is None:
        return 0
    span = (cout - cin) % 1440
    return int(setting_num("pay_break_default", 60)) if span > 360 else 0


def month_cards(staff_id, first, last):
    return get_db().execute("SELECT * FROM timecards WHERE staff_id=? AND date BETWEEN ? AND ? ORDER BY date, clock_in",
                            (staff_id, first.isoformat(), last.isoformat())).fetchall()


def month_summary(staff_id, first, last):
    """月の合計。yk_〜 は夜勤の分だけ（夜勤を1回いくらで払うときは、時間の計算から外すため）"""
    s = {"days": set(), "day_days": set(), "total": 0, "over": 0, "night": 0, "yakin": 0, "missing": 0,
         "yk_total": 0, "yk_over": 0, "yk_night": 0, "yk_cards": []}
    all_day, not_yk = {}, {}  # 日ごとの実働（全部／夜勤以外）
    for c in month_cards(staff_id, first, last):
        w = work_minutes(c)
        if w is None:
            s["missing"] += 1
            continue
        s["days"].add(c["date"])
        all_day[c["date"]] = all_day.get(c["date"], 0) + w["total"]
        for k in ("total", "night"):
            s[k] += w[k]
        if w["yakin"]:
            s["yakin"] += 1
            for k in ("total", "night"):
                s["yk_" + k] += w[k]
            s["yk_cards"].append((c, w))
        else:
            s["day_days"].add(c["date"])
            not_yk[c["date"]] = not_yk.get(c["date"], 0) + w["total"]
    # 残業は1日の合計で8時間をこえた分。夜勤を1回いくらで払うときは、夜勤の時間を外した残りで数える（over − yk_over）
    s["over"] = sum(max(0, t - 480) for t in all_day.values())
    s["yk_over"] = s["over"] - sum(max(0, t - 480) for t in not_yk.values())
    s["days"], s["day_days"] = len(s["days"]), len(s["day_days"])
    return s


def health_of(staff_id, d):
    return get_db().execute("SELECT * FROM health_checks WHERE staff_id=? AND date=? ORDER BY time", (staff_id, d)).fetchall()


def is_unwell(h):
    return (h["temp"] or 0) >= fever_line() or bool(h["symptoms"])


def _save_health(staff_id, form, username):
    temp = form.get("temp", type=finite_float)
    temp = round(temp, 1) if temp is not None else None
    symptoms = "、".join(s for s in SYMPTOMS if form.get(f"sym::{s}"))
    if temp is None and not symptoms and not form.get("health_note"):
        return None
    t = datetime.now()
    get_db().execute("INSERT INTO health_checks (staff_id, date, time, temp, symptoms, note, updated_by, updated_at) VALUES (?,?,?,?,?,?,?,?)",
                     (staff_id, t.date().isoformat(), t.strftime("%H:%M"), temp, symptoms, form.get("health_note") or "",
                      username, now()))
    return {"temp": temp, "symptoms": symptoms}


def punch(sid, action, form, username):
    """出勤・退勤・体温の記録。時刻はこのPCの時計で決める（本人は時刻を入れられない）。メッセージの一覧を返す"""
    db = get_db()
    t = datetime.now()
    card = open_card(sid)
    msgs = []
    if action == "in":
        st = _started(card) if card else None
        if card and st and t - st <= timedelta(minutes=forgot_minutes()):
            return [("error", f"すでに {card['clock_in']} に出勤しています。退勤のときは「退勤する」を押してください。")]
        if get_setting("pay_health_required", "1") == "1" and form.get("temp", type=finite_float) is None:
            return [("error", "出勤の前に体温を入れてください。")]
        # 退勤を押さないまま時間がたった打刻は、そのまま「退勤忘れ」として残し（管理者が直す）、新しく出勤する
        # 前回（いちばん新しい打刻）の退勤がないときだけ知らせる（もっと前の退勤忘れは、打刻の画面と管理者のホームに出る）
        old = get_db().execute("SELECT * FROM timecards WHERE staff_id=? ORDER BY date DESC, clock_in DESC, id DESC LIMIT 1",
                               (sid,)).fetchone()
        if old is not None and not old["clock_out"]:
            od = parse_date(old["date"])
            msgs.append(("warn", f"前回（{od.month}/{od.day} {old['clock_in'] or ''}〜）の退勤が押されていません。"
                                  "管理者に直してもらってください。" if od else
                                  "前回の退勤が押されていません。管理者に直してもらってください。"))
        h = _save_health(sid, form, username)
        db.execute("INSERT INTO timecards (staff_id, date, clock_in, note, updated_by, updated_at) VALUES (?,?,?,?,?,?)",
                   (sid, t.date().isoformat(), t.strftime("%H:%M"), form.get("note") or "", username, now()))
        log_event("clock_in", "timecards", sid, t.strftime("%H:%M"), username=username)
        if h and ((h["temp"] or 0) >= fever_line() or h["symptoms"]):
            msgs.append(("error", "体調がよくないようです。勤務の前に管理者に連絡してください。管理者のホームにもお知らせが出ます。"))
        msgs.append(("ok", f"{t:%H:%M} 出勤しました。今日もよろしくお願いします。"))
    elif action == "out":
        if not card:
            return [("error", "出勤の打刻が見つかりません。出勤を押し忘れたときや、長い時間がたったときは管理者に直してもらってください。")]
        out = t.strftime("%H:%M")
        brk = _break_min(form.get("break_min"))
        db.execute("UPDATE timecards SET clock_out=?, break_min=?, updated_by=?, updated_at=? WHERE id=?",
                   (out, default_break(card, out) if brk is None else brk, username, now(), card["id"]))
        log_event("clock_out", "timecards", sid, out, username=username)
        msgs.append(("ok", f"{out} 退勤しました。おつかれさまでした。"))
    elif action == "health":
        if _save_health(sid, form, username) is not None:
            msgs.append(("ok", "体温・体調を記録しました。"))
    return msgs


# ---------------------------------------------------------------- 事務所のPCの打刻画面（ログインしないで使う）
def _kiosk_people():
    """打刻の画面に出す人（職員の情報とつないだログイン。管理者は出さない：管理者はログインしてから打刻する）"""
    db = get_db()
    rows = []
    for s in db.execute("SELECT s.*, u.id AS uid, u.pin_hash, u.locked_until FROM staff s JOIN users u ON u.staff_id=s.id "
                        "WHERE u.active=1 AND u.role != 'admin' AND (s.status IS NULL OR s.status != '退職') ORDER BY s.kana, s.name"):
        card = open_card(s["id"])
        rows.append({"s": s, "card": card, "stale": is_stale(card), "pin": bool(s["pin_hash"])})
    return rows


KIOSK_MAX_FAILS = 5
KIOSK_LOCK_MINUTES = 15


def _kiosk_fail_state(user_id):
    """打刻のPINの失敗（ログインのパスワードとは別に数える）。(回数, この時刻までロック) を返す"""
    raw = get_setting(f"kiosk_fail:{user_id}", "") or ""
    count, _, until = raw.partition("|")
    try:
        count = int(count or 0)
    except ValueError:
        count = 0
    return count, until or None


def _kiosk_locked(user_id):
    _, until = _kiosk_fail_state(user_id)
    return bool(until) and until > now()


def _kiosk_fail(user_id):
    count, _ = _kiosk_fail_state(user_id)
    count += 1
    until = ""
    if count >= KIOSK_MAX_FAILS:
        until = (datetime.now() + timedelta(minutes=KIOSK_LOCK_MINUTES)).strftime("%Y-%m-%d %H:%M:%S")
        count = 0
    get_db().execute("INSERT OR REPLACE INTO settings (key, value) VALUES (?, ?)", (f"kiosk_fail:{user_id}", f"{count}|{until}"))


def _kiosk_reset(user_id):
    get_db().execute("DELETE FROM settings WHERE key=?", (f"kiosk_fail:{user_id}",))


def _timecard_off():
    from .customize import features_off

    return "timecard" in features_off()


@bp.route("/kiosk", methods=["GET", "POST"])
def kiosk():
    """名前を押して PIN → 出勤／退勤。押したらすぐ名前の一覧に戻る（ログインしたままにならない）。
    PINでログインできる端末として登録した事務所のPCでだけ開ける"""
    from werkzeug.security import check_password_hash

    from .auth import current_device

    if _timecard_off() or punch_mode() != "punch":
        # 「使う機能」でタイムカードを止めているとき・勤務表から確定するときは打刻できない
        if request.method == "POST":
            abort(403)
        return render_template("work_kiosk.html", off=True, device=None, people=[], person=None)
    if request.method == "GET":
        # 打刻の画面を開いたら、このPCでログインしたままの人は必ずログアウトにする（管理者の画面が後ろで開いたままにならない）
        for k in ("uid", "seen", "via", "csrf"):
            session.pop(k, None)
        g.user = None
    device = current_device()
    if device is None:
        return render_template("work_kiosk.html", device=None, people=[], person=None)
    db = get_db()
    if not session.get("kiosk_csrf"):
        import secrets

        session["kiosk_csrf"] = secrets.token_hex(16)
    sid = request.values.get("staff_id", type=db_int)
    person = next((p for p in _kiosk_people() if p["s"]["id"] == sid), None) if sid else None
    if request.method == "POST" and person:
        if request.form.get("_k") != session.get("kiosk_csrf"):
            abort(400)
        user = db.execute("SELECT * FROM users WHERE id=?", (person["s"]["uid"],)).fetchone()
        if _kiosk_locked(user["id"]):
            flash(f"PINを続けて間違えたため、{KIOSK_LOCK_MINUTES}分ほど打刻できません。管理者に連絡してください。", "error")
            log_event("login_failed", username=user["username"], detail=f"打刻のPIN・ロック中（{device['name']}）")
        elif not user["pin_hash"] or not check_password_hash(user["pin_hash"], request.form.get("pin", "")):
            # 打刻のPINの失敗は打刻の画面だけで数える（パスワードでのログインはロックしない）
            _kiosk_fail(user["id"])
            log_event("login_failed", username=user["username"], detail=f"打刻のPIN（{device['name']}）")
            flash("PINが違います。", "error")
            db.commit()
            return redirect(url_for("work.kiosk", staff_id=sid))
        else:
            _kiosk_reset(user["id"])
            db.execute("UPDATE devices SET last_used=? WHERE id=?", (now(), device["id"]))
            msgs = punch(sid, request.form.get("action"), request.form, user["username"])
            for cat, msg in msgs:
                flash(f"{person['s']['name']} さん：{msg}", cat)
            db.commit()
            if any(cat == "error" for cat, _ in msgs) and request.form.get("action") == "in" and (not person["card"] or person["stale"]):
                return redirect(url_for("work.kiosk", staff_id=sid))
        db.commit()
        return redirect(url_for("work.kiosk"))
    return render_template("work_kiosk.html", device=device, people=_kiosk_people(), person=person, SYMPTOMS=SYMPTOMS,
                           default_break=int(setting_num("pay_break_default", 60)), now_time=datetime.now(),
                           kcsrf=session["kiosk_csrf"], fever=fever_line())


# ---------------------------------------------------------------- 出勤・退勤（職員の画面）
@bp.route("/clock", methods=["GET", "POST"])
def clock():
    if punch_mode() != "punch":
        return redirect(url_for("work.confirm") if admin_view() else url_for("work.my_health"))
    db = get_db()
    sid = my_staff_id()
    staff = db.execute("SELECT * FROM staff WHERE id=?", (sid,)).fetchone() if sid else None
    card = open_card(sid) if sid else None
    if request.method == "POST":
        if not staff:
            abort(400)
        for cat, msg in punch(sid, request.form.get("action"), request.form, g.user["username"]):
            flash(msg, cat)
        db.commit()
        return redirect(url_for("work.clock"))
    first, last = parse_ym(None)
    cards = []
    if sid:
        for c in reversed(month_cards(sid, first, last)[-10:]):
            cards.append({"c": c, "w": work_minutes(c)})
    return render_template("work_clock.html", staff=staff, card=card, stale=is_stale(card), cards=cards, SYMPTOMS=SYMPTOMS, hm=hm,
                           forgotten=forgotten_cards(sid) if sid else [], max_hours=max_shift_minutes() // 60,
                           default_break=int(setting_num("pay_break_default", 60)), fever=fever_line(),
                           summary=month_summary(sid, first, last) if sid else None,
                           today_health=health_of(sid, date.today().isoformat()) if sid else [], now_time=datetime.now())


# ---------------------------------------------------------------- タイムカード（月ごと。管理者は直せる）
def _break_min(raw):
    """休憩（分）。読めない値は空、0〜1440分（1日）の間におさめる"""
    n = safe_int(raw, None)
    return None if n is None else clamp(n, 0, 24 * 60)


def _staff_list():
    return get_db().execute("SELECT * FROM staff WHERE status IS NULL OR status != '退職' ORDER BY kana, name").fetchall()


@bp.route("/timecards", methods=["GET", "POST"])
def timecards():
    db = get_db()
    first, last = parse_ym(request.values.get("ym"))
    ym = first.strftime("%Y-%m")
    if is_admin():
        require_admin()  # PINで入った管理者は、パスワードで確認してから
    admin = admin_view()
    staff_list = _staff_list() if admin else []
    sid = request.values.get("staff_id", type=db_int) if admin else my_staff_id()
    if admin and not sid and staff_list:
        sid = staff_list[0]["id"]
    staff = db.execute("SELECT * FROM staff WHERE id=?", (sid,)).fetchone() if sid else None
    if request.method == "POST":
        if not admin or not staff:
            abort(403)
        changed = 0
        kept = []  # 出勤の時刻が空・まちがいで、前の時刻のままにした日
        for c in month_cards(sid, first, last):
            if request.form.get(f"del_{c['id']}"):
                db.execute("DELETE FROM timecards WHERE id=?", (c["id"],))
                db.execute("DELETE FROM attend_days WHERE card_id=?", (c["id"],))  # 勤務の確定も「まだ」にもどす
                changed += 1
                continue
            cin = (request.form.get(f"in_{c['id']}") or "").strip()
            if _min(cin) is None:
                # 出勤の時刻は消せない（行を消すときは「消す」にチェック）。前の時刻のままにする
                cin = c["clock_in"]
                if (request.form.get(f"in_{c['id']}") or "").strip() != (c["clock_in"] or ""):
                    kept.append(c["date"][5:].replace("-", "/"))
            cout = (request.form.get(f"out_{c['id']}") or "").strip() or None
            if cout is not None and _min(cout) is None:
                cout = c["clock_out"]
            vals = (cin, cout, _break_min(request.form.get(f"br_{c['id']}")), request.form.get(f"note_{c['id']}") or "")
            if vals != (c["clock_in"], c["clock_out"], c["break_min"], c["note"] or ""):
                db.execute("UPDATE timecards SET clock_in=?, clock_out=?, break_min=?, note=?, updated_by=?, updated_at=? WHERE id=?",
                           vals + (g.user["username"], now(), c["id"]))
                changed += 1
        for d in month_days(first, last):
            cin = (request.form.get(f"in_new_{d.day}") or "").strip()
            if _min(cin) is not None:
                db.execute("INSERT INTO timecards (staff_id, date, clock_in, clock_out, break_min, note, updated_by, updated_at)"
                           " VALUES (?,?,?,?,?,?,?,?)", (sid, d.isoformat(), cin, request.form.get(f"out_new_{d.day}") or None,
                                                        _break_min(request.form.get(f"br_new_{d.day}")), request.form.get(f"note_new_{d.day}") or "",
                                                        g.user["username"], now()))
                changed += 1
        if changed:
            log_event("timecard_edit", "timecards", sid, f"{ym} {changed}件")
        db.commit()
        flash(f"{staff['name']} さんの{first:%Y年%m月}のタイムカードを保存しました（{changed}件）。", "ok")
        if kept:
            flash(f"{'、'.join(kept)} は出勤の時刻が空か正しくないため、前の時刻のままにしました。"
                  "その日の打刻をなくすときは、右の「消す」にチェックしてください。", "error")
        return redirect(url_for("work.timecards", ym=ym, staff_id=sid))
    rows = []
    if staff:
        by_day = {}
        for x in day_cards(month_cards(sid, first, last)):
            by_day.setdefault(x["c"]["date"], []).append(x)
        temps = {}
        for h in db.execute("SELECT * FROM health_checks WHERE staff_id=? AND date BETWEEN ? AND ? ORDER BY time",
                            (sid, first.isoformat(), last.isoformat())):
            temps.setdefault(h["date"], h)
        plan = planned_shifts(sid, first, last)
        tmap = shift_types_map()
        diffs = shift_differences(sid, first, last)
        for d in month_days(first, last):
            cs = by_day.get(d.isoformat(), [])
            code = plan.get(d.isoformat(), "")
            rows.append({"d": d, "cards": cs, "health": temps.get(d.isoformat()),
                         "plan": code, "plan_t": tmap.get(code), "diffs": diffs.get(d.isoformat(), [])})
    return render_template("work_timecards.html", staff=staff, staff_list=staff_list, rows=rows, ym=ym, first=first, WEEK=WEEK,
                           hm=hm, summary=month_summary(sid, first, last) if staff else None, admin=admin, fever=fever_line(),
                           night_label=night_label())


@bp.route("/timecards.xlsx")
@admin_required
def timecards_export():
    first, last = parse_ym(request.args.get("ym"))
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for s in _staff_list():
        cards = month_cards(s["id"], first, last)
        if not cards:
            continue
        rows = []
        for x in day_cards(cards):
            c, w = x["c"], x["w"] or {}
            d = parse_date(c["date"])
            rows.append([f"{d.month}/{d.day}", WEEK[d.weekday()], c["clock_in"] or "", c["clock_out"] or "", c["break_min"] or 0,
                         hm(w.get("total")), hm(w.get("over")), hm(w.get("night")), "夜勤" if w.get("yakin") else "", c["note"] or ""])
        sm = month_summary(s["id"], first, last)
        rows.append(["合計", "", "", "", "", hm(sm["total"]), hm(sm["over"]), hm(sm["night"]), f"{sm['yakin']}回", f"出勤 {sm['days']}日"])
        excel.add_table(wb.create_sheet(excel.safe_sheet_title(s["name"])), f"出勤簿　{s['name']}　{first:%Y年%m月}",
                        ["日", "曜日", "出勤", "退勤", "休憩（分）", "実働", "残業", "深夜", "夜勤", "備考"], rows,
                        widths=[7, 5, 8, 8, 9, 8, 8, 8, 7, 24])
    if not wb.sheetnames:
        wb.create_sheet("なし")
    return excel.send_workbook(wb, f"出勤簿_{first:%Y%m}")


# ---------------------------------------------------------------- 体温（健康チェック）
@bp.route("/health")
def health():
    db = get_db()
    if is_admin():
        require_admin()  # PINで入った管理者は、パスワードで確認してから全員の体温を見る
    if not admin_view():
        sid = my_staff_id()
        rows = db.execute("SELECT * FROM health_checks WHERE staff_id=? ORDER BY date DESC, time DESC LIMIT 60", (sid,)).fetchall() if sid else []
        return render_template("work_health.html", mine=True, rows=rows, fever=fever_line(), day=date.today())
    day = parse_date(request.args.get("date"), date.today())
    checks = {}
    for h in db.execute("SELECT * FROM health_checks WHERE date=? ORDER BY time", (day.isoformat(),)):
        checks.setdefault(h["staff_id"], []).append(h)
    on_duty = {c["staff_id"] for c in db.execute("SELECT staff_id FROM timecards WHERE date=?", (day.isoformat(),))}
    rows = [{"s": s, "checks": checks.get(s["id"], []), "on": s["id"] in on_duty} for s in _staff_list()]
    rows.sort(key=lambda x: (not any(is_unwell(h) for h in x["checks"]), not x["on"]))
    return render_template("work_health.html", mine=False, rows=rows, fever=fever_line(), day=day, is_unwell=is_unwell,
                           prev=(day - timedelta(days=1)).isoformat(), next=(day + timedelta(days=1)).isoformat())


def today_status():
    """管理者のホーム用：出勤中の人・体調がよくない人・打刻漏れ"""
    db = get_db()
    today = date.today().isoformat()
    names = {s["id"]: s["name"] for s in db.execute("SELECT id, name FROM staff")}
    forgot = {c["id"] for c in forgotten_cards(days=3)}
    working = [names.get(c["staff_id"], "?") + f"（{c['clock_in']}〜）" for c in db.execute(
        "SELECT * FROM timecards WHERE clock_out IS NULL AND date >= ? ORDER BY clock_in",
        ((date.today() - timedelta(days=2)).isoformat(),)) if c["id"] not in forgot]
    unwell = [(names.get(h["staff_id"], "?"), h) for h in db.execute("SELECT * FROM health_checks WHERE date=? ORDER BY time", (today,))
              if is_unwell(h)]
    confirm = None
    if punch_mode() == "shift":
        rows, _ = confirm_rows(date.today())
        confirm = {"total": len(rows), "done": sum(1 for r in rows if r["a"])}
    return {"working": working, "unwell": unwell, "confirm": confirm}


# ---------------------------------------------------------------- 勤務表（管理者が組む）とのつき合わせ
def shift_types_map():
    return {t["code"]: t for t in get_db().execute("SELECT * FROM shift_types")}


def planned_shifts(staff_id, first, last):
    return {s["date"]: s["code"] for s in get_db().execute(
        "SELECT * FROM shifts WHERE staff_id=? AND date BETWEEN ? AND ?", (staff_id, first.isoformat(), last.isoformat()))}


def _is_work_type(t):
    return bool(t and t["start"] and _min(t["start"]) is not None)


def shift_differences(staff_id, first, last):
    """勤務表と打刻のちがい（今日より前の日だけ）。{日付: [(種類, 説明)]}"""
    grace = int(setting_num("pay_late_grace", 10))
    tmap = shift_types_map()
    plan = planned_shifts(staff_id, first, last)
    cards = {}
    for c in month_cards(staff_id, first, last):
        cards.setdefault(c["date"], []).append(c)
    out = {}
    today = date.today()
    for d in month_days(first, last):
        if d >= today:
            break
        key = d.isoformat()
        t = tmap.get(plan.get(key, ""))
        cs = cards.get(key, [])
        diffs = []
        word = "打刻" if punch_mode() == "punch" else "出勤簿"
        if _is_work_type(t) and not cs:
            diffs.append(("打刻なし" if word == "打刻" else "確定なし", f"勤務表は「{t['name']}」ですが{word}がありません"))
        elif cs and not _is_work_type(t):
            diffs.append(("予定外", f"勤務表は「{t['name'] if t else '空欄'}」ですが {cs[0]['clock_in'] or '時刻なし'} に{word}があります"))
        elif cs and t:
            first_in = min((_min(c["clock_in"]) for c in cs if _min(c["clock_in"]) is not None), default=None)
            ps, pe = _min(t["start"]), _min(t["end"])
            if first_in is not None and ps is not None and first_in > ps + grace:
                diffs.append(("遅刻", f"予定 {t['start']} → 出勤 {first_in // 60}:{first_in % 60:02d}"))
            last = cs[-1]
            w = work_minutes(last)
            if pe is not None and last["clock_out"] and w is not None:
                end_plan = pe + (1440 if pe <= ps else 0)
                end_real = _min(last["clock_out"]) + (1440 if _min(last["clock_out"]) <= _min(last["clock_in"]) else 0)
                if end_real < end_plan - grace:
                    diffs.append(("早退", f"予定 {t['end']} → 退勤 {last['clock_out']}"))
            yk = any((work_minutes(c) or {}).get("yakin") for c in cs)
            if bool(t["night"]) != yk:
                diffs.append(("夜勤のちがい", "勤務表は夜勤ですが、打刻は夜勤になっていません" if t["night"]
                              else "勤務表は夜勤ではありませんが、打刻が夜勤（日をまたぐ）になっています"))
        if diffs:
            out[key] = diffs
    return out


def my_upcoming_shifts(staff_id, days=14):
    tmap = shift_types_map()
    start = date.today()
    end = start + timedelta(days=days - 1)
    plan = {}
    for s in get_db().execute("SELECT * FROM shifts WHERE staff_id=? AND date BETWEEN ? AND ?",
                              (staff_id, start.isoformat(), end.isoformat())):
        plan[s["date"]] = s["code"]
    rows = []
    for i in range(days):
        d = start + timedelta(days=i)
        t = tmap.get(plan.get(d.isoformat(), ""))
        rows.append({"d": d, "w": WEEK[d.weekday()], "code": plan.get(d.isoformat(), ""), "t": t})
    return rows


@bp.route("/my-shift")
def my_shift():
    """職員：自分の勤務表（見るだけ。勤務表は管理者が組む）"""
    first, last = parse_ym(request.args.get("ym"))
    sid = my_staff_id()
    tmap = shift_types_map()
    plan = planned_shifts(sid, first, last) if sid else {}
    rows = [{"d": d, "w": WEEK[d.weekday()], "code": plan.get(d.isoformat(), ""), "t": tmap.get(plan.get(d.isoformat(), ""))}
            for d in month_days(first, last)]
    hours = sum((r["t"]["hours"] or 0) for r in rows if r["t"])
    from .leave import my_balance

    return render_template("work_my_shift.html", rows=rows, first=first, ym=first.strftime("%Y-%m"), linked=bool(sid),
                           leave=my_balance(sid) if sid else None,
                           hours=hours, today=date.today(), types=list(tmap.values()))


# ---------------------------------------------------------------- 勤務の確定（勤務表から。打刻のかわり）
CONFIRM_NOTE = "勤務表で確定"


def off_kinds():
    from .leave import LEAVE_CODE

    return {"休": "休み", LEAVE_CODE: "有給"}


def shift_break(t):
    """勤務の種類の休憩（分）：開始〜終了の長さから「勤務時間（休憩を除く）」を引いた分。
    勤務時間が入っていない（0の宿直など）ときは、ふつうの休憩（6時間をこえたら設定の分）"""
    s, e = _min(t["start"]), _min(t["end"])
    if s is None or e is None:
        return 0
    span = (e - s) % 1440 or 1440
    hours = t["hours"]
    if hours:
        return clamp(span - round(float(hours) * 60), 0, span)
    return int(setting_num("pay_break_default", 60)) if span > 360 else 0


def attend_of(staff_id, d):
    return get_db().execute("SELECT * FROM attend_days WHERE staff_id=? AND date=?", (staff_id, d)).fetchone()


def _set_shift(staff_id, d, code, username):
    db = get_db()
    if code:
        db.execute("INSERT OR REPLACE INTO shifts (staff_id, date, code, updated_by, updated_at) VALUES (?,?,?,?,?)",
                   (staff_id, d, code, username, now()))
    else:
        db.execute("DELETE FROM shifts WHERE staff_id=? AND date=?", (staff_id, d))


def undo_confirm(staff_id, d, username):
    """確定を取り消す：作ったタイムカードを消し、休み・有給で直した勤務表を元にもどす"""
    db = get_db()
    a = attend_of(staff_id, d)
    if not a:
        return False
    if a["card_id"]:
        db.execute("DELETE FROM timecards WHERE id=?", (a["card_id"],))
    if a["status"] != "work" or a["planned_code"] != (planned_shifts_on(staff_id, d) or ""):
        _set_shift(staff_id, d, a["planned_code"] or "", username)
    db.execute("DELETE FROM attend_days WHERE staff_id=? AND date=?", (staff_id, d))
    return True


def planned_shifts_on(staff_id, d):
    r = get_db().execute("SELECT code FROM shifts WHERE staff_id=? AND date=?", (staff_id, d)).fetchone()
    return r["code"] if r else None


def confirm_work(staff_id, d, start, end, brk, username, code=None):
    """出勤で確定：その時間のタイムカードを作る。code（勤務表にない日に足すとき）は勤務表にも入れる"""
    db = get_db()
    undo_confirm(staff_id, d, username)
    planned = planned_shifts_on(staff_id, d) or ""
    if code and code != planned:
        _set_shift(staff_id, d, code, username)
    cur = db.execute("INSERT INTO timecards (staff_id, date, clock_in, clock_out, break_min, note, updated_by, updated_at)"
                     " VALUES (?,?,?,?,?,?,?,?)", (staff_id, d, start, end, brk, CONFIRM_NOTE, username, now()))
    db.execute("INSERT INTO attend_days (staff_id, date, status, planned_code, card_id, confirmed_by, confirmed_at) VALUES (?,?,?,?,?,?,?)",
               (staff_id, d, "work", planned, cur.lastrowid, username, now()))


def confirm_off(staff_id, d, kind, username):
    """休み・有給で確定：勤務表をそのように直す（元の予定は覚えておき、取り消すと元にもどす）"""
    db = get_db()
    undo_confirm(staff_id, d, username)
    planned = planned_shifts_on(staff_id, d) or ""
    _set_shift(staff_id, d, kind, username)
    db.execute("INSERT INTO attend_days (staff_id, date, status, planned_code, card_id, confirmed_by, confirmed_at) VALUES (?,?,?,?,?,?,?)",
               (staff_id, d, "paid" if kind == off_kinds_paid() else "off", planned, None, username, now()))


def off_kinds_paid():
    from .leave import LEAVE_CODE

    return LEAVE_CODE


def confirm_rows(d):
    """その日の勤務表と確定の状態。勤務の日（開始の時間がある種類）と、確定した人を出す"""
    db = get_db()
    tmap = shift_types_map()
    ds = d.isoformat()
    plan = {r["staff_id"]: r["code"] for r in db.execute("SELECT staff_id, code FROM shifts WHERE date=?", (ds,))}
    att = {r["staff_id"]: r for r in db.execute("SELECT * FROM attend_days WHERE date=?", (ds,))}
    cards = {r["id"]: r for r in db.execute("SELECT * FROM timecards WHERE date=?", (ds,))}
    temps = {}
    for h in db.execute("SELECT * FROM health_checks WHERE date=? ORDER BY time", (ds,)):
        temps.setdefault(h["staff_id"], h)
    rows = []
    for s in _staff_list():
        a = att.get(s["id"])
        code = plan.get(s["id"], "")
        t = tmap.get(code)
        orig = tmap.get(a["planned_code"]) if a else t  # 休み・有給で確定した日は、元の予定
        if not a and not _is_work_type(t):
            continue
        card = cards.get(a["card_id"]) if a and a["card_id"] else None
        rows.append({"s": s, "code": code, "t": t, "orig": orig, "a": a, "card": card, "health": temps.get(s["id"]),
                     "brk": shift_break(orig) if _is_work_type(orig) else 0})
    others = [s for s in _staff_list() if s["id"] not in {r["s"]["id"] for r in rows}]
    return rows, others


def unconfirmed_days(days=31):
    """今日までで、勤務の日（勤務表）なのに確定していない日：{日付: [職員の名前]}。使い始めた日より前は見ない"""
    from .compliance import before_start

    db = get_db()
    today = date.today()
    since = (today - timedelta(days=days)).isoformat()
    work_codes = [c for c, t in shift_types_map().items() if _is_work_type(t)]
    if not work_codes:
        return {}
    q = ",".join("?" * len(work_codes))
    out = {}
    for r in db.execute(f"SELECT sh.date, st.name FROM shifts sh JOIN staff st ON st.id=sh.staff_id "
                        f"WHERE sh.code IN ({q}) AND sh.date BETWEEN ? AND ? AND (st.status IS NULL OR st.status != '退職') "
                        "AND NOT EXISTS (SELECT 1 FROM attend_days a WHERE a.staff_id=sh.staff_id AND a.date=sh.date) "
                        "ORDER BY sh.date, st.kana", work_codes + [since, today.isoformat()]):
        if before_start(parse_date(r["date"])):
            continue
        out.setdefault(r["date"], []).append(r["name"])
    return out


def _hhmm(raw, default):
    v = (raw or "").strip()
    return v if _min(v) is not None and len(v) <= 5 else default


@bp.route("/confirm", methods=["GET", "POST"])
@admin_required
def confirm():
    """勤務の確定：勤務表どおりに働いたら「出勤」、休んだら「休み」「有給」。時間がちがうときは直してから「出勤」"""
    db = get_db()
    today = date.today()
    d = parse_date(request.values.get("date"), today)
    if d > today:
        d = today
    ds = d.isoformat()
    if request.method == "POST":
        action = request.form.get("action", "")
        sid = request.form.get("staff_id", type=db_int)
        staff = db.execute("SELECT * FROM staff WHERE id=?", (sid,)).fetchone() if sid else None
        tmap = shift_types_map()
        user = g.user["username"]
        msg = None
        if action == "all":
            rows, _ = confirm_rows(d)
            n = 0
            for r in rows:
                if not r["a"] and _is_work_type(r["t"]):
                    confirm_work(r["s"]["id"], ds, r["t"]["start"], r["t"]["end"], shift_break(r["t"]), user)
                    n += 1
            msg = f"{n}人を勤務表どおりに確定しました。"
            log_event("attend_confirm", "attend_days", None, f"{ds} まとめて {n}人")
        elif staff is None:
            abort(400)
        elif action in ("work", "add"):
            code = request.form.get("code", "") if action == "add" else None
            if action == "add":
                t = tmap.get(code or "")
            else:  # 休み・有給で確定したあとに出勤へ直すときは、元の予定の時間
                a = attend_of(sid, ds)
                t = tmap.get((a["planned_code"] if a else planned_shifts_on(sid, ds)) or "")
            if action == "add" and not _is_work_type(t):
                flash("勤務の種類を選んでください。", "error")
                return redirect(url_for("work.confirm", date=ds))
            start = _hhmm(request.form.get("start"), t["start"] if t else None)
            end = _hhmm(request.form.get("end"), t["end"] if t else None)
            if start is None or end is None:
                flash(f"{staff['name']} さん：始まりと終わりの時間を入れてください。", "error")
                return redirect(url_for("work.confirm", date=ds))
            brk = _break_min(request.form.get("break_min"))
            if brk is None:
                brk = shift_break(t) if _is_work_type(t) else default_break({"clock_in": start}, end)
            confirm_work(sid, ds, start, end, brk, user, code=code)
            temp = request.form.get("temp", type=finite_float)
            if temp is not None:
                db.execute("INSERT INTO health_checks (staff_id, date, time, temp, symptoms, note, updated_by, updated_at) VALUES (?,?,?,?,?,?,?,?)",
                           (sid, ds, start, round(temp, 1), "", "勤務の確定で記録", user, now()))
            msg = f"{staff['name']} さん：{start}〜{end} の出勤で確定しました。"
            log_event("attend_confirm", "attend_days", sid, f"{ds} 出勤 {start}〜{end}")
        elif action.startswith("off:"):
            kind = action[4:]
            kinds = off_kinds()
            if kind not in kinds:
                abort(400)
            if kind not in tmap:
                flash(f"勤務の種類に「{kind}」がありません。「勤務の種類」に追加してください。", "error")
                return redirect(url_for("work.confirm", date=ds))
            confirm_off(sid, ds, kind, user)
            msg = f"{staff['name']} さん：{kinds[kind]}で確定しました（勤務表も「{kind}」にしました）。"
            log_event("attend_confirm", "attend_days", sid, f"{ds} {kinds[kind]}")
        elif action == "undo":
            if undo_confirm(sid, ds, user):
                msg = f"{staff['name']} さんの確定を取り消しました。"
                log_event("attend_undo", "attend_days", sid, ds)
        else:
            abort(400)
        db.commit()
        if msg:
            flash(msg, "ok")
        return redirect(url_for("work.confirm", date=ds) + (f"#s{sid}" if sid else ""))
    rows, others = confirm_rows(d)
    work_types = [t for t in shift_types_map().values() if _is_work_type(t)]
    left = unconfirmed_days()
    left.pop(ds, None)
    return render_template("work_confirm.html", d=d, rows=rows, others=others, work_types=work_types, WEEK=WEEK, hm=hm,
                           kinds=off_kinds(), left=left, today=today, fever=fever_line(), work_minutes=work_minutes,
                           prev=(d - timedelta(days=1)).isoformat(), next=(d + timedelta(days=1)).isoformat() if d < today else None)


@bp.route("/my-health", methods=["GET", "POST"])
def my_health():
    """職員：自分の体温・体調を記録する（勤務表から確定するときの、打刻のかわり）"""
    sid = my_staff_id()
    db = get_db()
    if request.method == "POST":
        if not sid:
            abort(400)
        if _save_health(sid, request.form, g.user["username"]) is None:
            flash("体温を入れてください。", "error")
        else:
            h = health_of(sid, date.today().isoformat())[-1]
            if is_unwell(h):
                flash("体調がよくないようです。勤務の前に管理者に連絡してください。管理者のホームにもお知らせが出ます。", "error")
            else:
                flash("体温・体調を記録しました。", "ok")
            log_event("health", "health_checks", sid)
        db.commit()
        return redirect(url_for("work.my_health"))
    rows = db.execute("SELECT * FROM health_checks WHERE staff_id=? ORDER BY date DESC, time DESC LIMIT 60", (sid,)).fetchall() if sid else []
    return render_template("work_my_health.html", rows=rows, linked=bool(sid), SYMPTOMS=SYMPTOMS, fever=fever_line(),
                           today_health=health_of(sid, date.today().isoformat()) if sid else [])
